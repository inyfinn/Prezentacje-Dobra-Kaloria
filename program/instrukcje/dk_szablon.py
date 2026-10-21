"""Złóż prezentację Dobra Kaloria ze slajdów pełnego szablonu (python-pptx >= 0.6.21).

python dk_szablon.py "DK - szablon prezentacji - pełen.pptx" wynik.pptx 4,17,62,62,65
-> wynik ma slajdy szablonu t04, t17, t62 (dwa razy), t65 w tej kolejności; reszta usunięta.
"""
import copy
import os
import re
import sys

from pptx import Presentation

R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
SKIP = ("notesSlide", "slideLayout")


def duplikuj(prs, src):
    """Kopia slajdu na końcu pliku: tło, kształty, obrazy, linki (rId przemapowane)."""
    new = prs.slides.add_slide(src.slide_layout)
    mapa = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype.endswith(SKIP):
            continue
        if rel.is_external:
            mapa[rid] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            mapa[rid] = new.part.relate_to(rel.target_part, rel.reltype)
    zrodlo = copy.deepcopy(src._element.cSld)
    for el in zrodlo.iter():
        for k, v in list(el.attrib.items()):
            if k.startswith(R) and v in mapa:
                el.set(k, mapa[v])
    # Podmieniamy dzieci spTree, nie sam element: python-pptx trzyma do niego odwołanie.
    tree = new.shapes._spTree
    for ch in list(tree):
        tree.remove(ch)
    for ch in list(zrodlo.spTree):
        tree.append(ch)
    if zrodlo.bg is not None:  # tło slajdu
        tree.addprevious(zrodlo.bg)
    return new


def zloz(prs, numery):
    """Zostaw tylko slajdy o numerach z listy (1 = pierwszy), w tej kolejności, z powtórzeniami."""
    src = list(prs.slides)
    lst = prs.slides._sldIdLst
    stare = list(lst)
    uzyte, kolejnosc = set(), []
    for n in numery:
        s = src[n - 1]
        if n in uzyte:
            duplikuj(prs, s)
            kolejnosc.append(lst[-1])
        else:
            uzyte.add(n)
            kolejnosc.append(stare[n - 1])
    for sid in stare:
        if sid not in kolejnosc:
            prs.part.drop_rel(sid.rId)
    for sid in list(lst):
        lst.remove(sid)
    for sid in kolejnosc:
        lst.append(sid)
    for i, (slide, n) in enumerate(zip(prs.slides, numery), 1):
        slide._element.attrib.pop("show", None)  # ukryte slajdy instrukcji -> zwykłe
        usun_instrukcje(prs, slide)
        if slide.has_notes_slide:  # notatki szablonu to opis slajdu, nie treść dla prelegenta
            slide.notes_slide.notes_text_frame.text = ""
        for sh in slide.shapes:  # numer strony w rogu = numer w szablonie -> nowy numer
            if sh.has_text_frame and sh.text_frame.text.strip() == f"{n:02d}" and sh.top > prs.slide_height * 0.8:
                wpisz(sh, f"{i:02d}")


def usun_instrukcje(prs, slide):
    """Usuń ramki „SLAJD NN - do czego służy” leżące poza kadrem."""
    W = prs.slide_width
    for sh in list(slide.shapes):
        if sh.left is not None and (sh.left >= W or sh.left + sh.width <= 0):
            sh._element.getparent().remove(sh._element)


def wpisz(shape_or_tf, *akapity):
    """Podmień tekst, zachowując format pierwszego fragmentu każdego akapitu."""
    tf = getattr(shape_or_tf, "text_frame", shape_or_tf)
    while len(tf.paragraphs) < len(akapity):  # brakuje akapitów: kopiuj ostatni
        tf.paragraphs[-1]._p.addnext(copy.deepcopy(tf.paragraphs[-1]._p))
    for i, p in enumerate(list(tf.paragraphs)):
        if i >= len(akapity):
            p._p.getparent().remove(p._p)
        elif p.runs:
            p.runs[0].text = akapity[i]
            for r in p.runs[1:]:
                r._r.getparent().remove(r._r)
        else:
            p.text = akapity[i]


def ksztalty(shapes):
    """Wszystkie kształty slajdu, także w grupach. Nazwy mówią, co to jest (np. 'Packshot - ...')."""
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:  # grupa
            yield from ksztalty(sh.shapes)


def obraz(slide, pic, plik):
    """Podmień obrazek w ramce szablonu; nowy mieści się w starej ramce w całości (bez kadrowania).
    Przezroczyste marginesy PNG (packshoty) są obcinane, żeby paczka nie wyszła mała."""
    import io
    from PIL import Image

    im = Image.open(plik)
    if im.mode in ("RGBA", "LA", "P"):
        im = im.convert("RGBA")
        bbox = im.getchannel("A").getbbox()
        if bbox and bbox != (0, 0) + im.size:
            im = im.crop(bbox)
    buf = io.BytesIO()
    im.save(buf, "PNG")
    buf.seek(0)
    _, rid = slide.part.get_or_add_image_part(buf)
    pic._element.blipFill.blip.rEmbed = rid
    kadr = pic._element.blipFill.find("{http://schemas.openxmlformats.org/drawingml/2006/main}srcRect")
    if kadr is not None:  # zdejmij kadrowanie starego obrazka
        pic._element.blipFill.remove(kadr)
    w, h = im.size
    L, T, W, H = pic.left, pic.top, pic.width, pic.height
    s = min(W / w, H / h)
    pic.width, pic.height = int(w * s), int(h * s)
    pic.left, pic.top = L + (W - pic.width) // 2, T + (H - pic.height) // 2


def pola(shapes, poczatek):
    """Pola tekstowe, których tekst zaczyna się od `poczatek` (np. "[Claim", "00%"), w kolejności czytania
    (wiersz po wierszu, od lewej). Twarde spacje i podziały linii szablonu traktuje jak zwykłe spacje."""
    norm = lambda s: " ".join(s.replace("\u00a0", " ").split())
    out = [sh for sh in ksztalty(shapes) if sh.has_text_frame and norm(sh.text_frame.text).startswith(norm(poczatek))]
    return sorted(out, key=lambda sh: (round(sh.top / 300000), sh.left))


def twarde(tekst):
    """Polska typografia: twarda spacja po jednoliterowych słowach (w, z, i, a, o, u), żeby nie zostały na końcu linii."""
    return re.sub(r"(?<!\S)([wWzZiIaAoOuU]) ", "\\1\u00a0", tekst)


A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
PLIKI_CZCIONEK = {"mindset": "Mindset.otf", "mindset slim": "MindsetSlim.otf",
                  "lato": "Lato-Regular.ttf", "lato bold": "Lato-Bold.ttf"}


def _czcionka(fonty, rp, size):
    """Czcionka PIL dla fragmentu tekstu: +mj-lt = Mindset (nagłówki), +mn-lt = Lato (tekst)."""
    from PIL import ImageFont

    lat = rp.find(A + "latin") if rp is not None else None
    tf = (lat.get("typeface") if lat is not None else "+mn-lt") or "+mn-lt"
    nazwa = {"+mj-lt": "mindset", "+mn-lt": "lato"}.get(tf, tf.lower())
    if nazwa == "lato" and rp is not None and rp.get("b") == "1":
        nazwa = "lato bold"
    plik = PLIKI_CZCIONEK.get(nazwa, "Lato-Regular.ttf")
    return ImageFont.truetype(os.path.join(fonty, plik), max(1, round(size * 4))), nazwa  # x4: dokładność


def _rozmiar(rp, p, domyslny):
    for el in (rp, p._p.find(A + "pPr/" + A + "defRPr") if p._p.find(A + "pPr") is not None else None):
        if el is not None and el.get("sz"):
            return int(el.get("sz")) / 100
    return domyslny


def sprawdz(prs, fonty):
    """Czy tekst mieści się w polach — liczone PRAWDZIWYMI czcionkami (Mindset, Lato), przy piśmie 100%.

    PowerPoint nie zmniejsza pisma sam (autodopasowanie bez zapisanej skali = 100%), a LibreOffice zmniejsza —
    dlatego podgląd z LibreOffice potrafi kłamać. Ta funkcja nie zależy od renderu.
    `fonty` = folder z Mindset.otf, MindsetSlim.otf, Lato-Regular.ttf, Lato-Bold.ttf.
    Zwraca listę usterek (slajd, pole, opis); pusta lista = wszystko się mieści.
    """
    usterki = []
    for nr, slide in enumerate(prs.slides, 1):
        for sh in ksztalty(slide.shapes):
            if not sh.has_text_frame or not sh.text_frame.text.strip() or sh.width is None:
                continue
            bp = sh.text_frame._txBody.bodyPr
            ins = lambda k, d=91440: int(bp.get(k, d))
            szer = (sh.width - ins("lIns") - ins("rIns")) / 12700  # pt
            wys = (sh.height - ins("tIns", 45720) - ins("bIns", 45720)) / 12700
            if bp.get("wrap") == "none":
                continue
            potrzeba, nazwa_pola = 0.0, sh.text_frame.text.strip().replace("\n", " / ")[:40]
            for p in sh.text_frame.paragraphs:
                runs = p.runs
                if not runs:
                    potrzeba += 0.6 * 18
                    continue
                rp = runs[0]._r.find(A + "rPr")
                size = _rozmiar(rp, p, 18)
                font, nazwa = _czcionka(fonty, rp, size)
                spc = int(rp.get("spc", 0)) / 100 if rp is not None else 0  # rozstrzelenie liter, pt
                asc, desc = font.getmetrics()
                ln = p._p.find(A + "pPr/" + A + "lnSpc/" + A + "spcPct")
                proc = int(ln.get("val")) / 100000 if ln is not None else 1.0
                linia = (asc + desc) / 4 * proc
                tekst = "".join(r.text for r in runs)
                for kawalek in tekst.split("\v"):
                    slowa, linie, biez = kawalek.split(" "), 1, ""
                    for s in slowa:  # twarda spacja   nie łamie — „2 800” to jedno słowo
                        w = (font.getlength(s) / 4) + spc * len(s)
                        if w > szer + 0.5:
                            usterki.append((nr, nazwa_pola, "słowo „%s” szersze niż pole (%.0f > %.0f pt) — złamie się w środku" % (s, w, szer)))
                        proba = (biez + " " + s).strip()
                        if (font.getlength(proba) / 4) + spc * len(proba) > szer and biez:
                            linie, biez = linie + 1, s
                        else:
                            biez = proba
                    potrzeba += linie * linia
            if potrzeba <= wys * 1.05 + 2:  # 5% zapasu na różnice silników
                continue
            kolizja = _kolizja(prs, slide, sh, (potrzeba - wys) * 12700, bp.get("anchor", "t"))
            if kolizja:
                usterki.append((nr, nazwa_pola, "tekst wymaga %.0f pt wysokości, pole ma %.0f pt — %s"
                                % (potrzeba, wys, kolizja)))
    return usterki


def _kolizja(prs, slide, sh, nadmiar, anchor):
    """Czy tekst rosnący o `nadmiar` (EMU) wejdzie na inny element, wyjdzie poza swoją kartę albo poza slajd."""
    if anchor == "ctr":
        gora, dol = sh.top - nadmiar / 2, sh.top + sh.height + nadmiar / 2
    elif anchor == "b":
        gora, dol = sh.top - nadmiar, sh.top + sh.height
    else:
        gora, dol = sh.top, sh.top + sh.height + nadmiar
    if gora < 0 or dol > prs.slide_height:
        return "wyjdzie poza slajd"
    for o in ksztalty(slide.shapes):
        if o._element is sh._element or o.shape_type == 6 or o.width is None or o.left >= prs.slide_width:
            continue
        if min(o.left + o.width, sh.left + sh.width) - max(o.left, sh.left) <= 12700 * 4:
            continue  # nie leżą nad sobą
        o_dol = o.top + o.height
        if o.left <= sh.left and o.left + o.width >= sh.left + sh.width and o.top <= sh.top and o_dol >= sh.top + sh.height:
            if dol > o_dol or gora < o.top:
                return "wyjdzie poza swoją kartę / pas"
            continue
        if o.top >= sh.top + sh.height - 12700 and o.top < dol:
            return "wejdzie na element poniżej („%s”)" % (o.text_frame.text.strip()[:25] if o.has_text_frame else o.name[:25])
        if o_dol <= sh.top + 12700 and o_dol > gora:
            return "wejdzie na element powyżej"
    return None


def podglad(prs_plik, wyjscie):
    """Kopia do renderu w LibreOffice bez autodopasowania, żeby podgląd pokazał to, co PowerPoint (pismo 100%)."""
    import zipfile

    with zipfile.ZipFile(prs_plik) as zin, zipfile.ZipFile(wyjscie, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            dane = zin.read(it.filename)
            if it.filename.startswith("ppt/slides/slide") and it.filename.endswith(".xml"):
                dane = re.sub(rb"<a:normAutofit[^>]*?(/>|>.*?</a:normAutofit>)", b"<a:noAutofit/>", dane, flags=re.S)
            zout.writestr(it, dane)


def nawiasy(prs):
    """Lista [podpowiedzi] zostawionych w pliku: (slajd, nazwa kształtu, tekst)."""
    out = []
    for i, s in enumerate(prs.slides, 1):
        for sh in ksztalty(s.shapes):
            if sh.has_text_frame and re.search(r"\[[^\]]+\]", sh.text_frame.text):
                out.append((i, sh.name, sh.text_frame.text))
    return out


if __name__ == "__main__":
    prs = Presentation(sys.argv[1])
    zloz(prs, [int(x) for x in sys.argv[3].split(",")])
    prs.save(sys.argv[2])
    for row in nawiasy(prs):
        print(*row, sep=" | ")
