# -*- coding: utf-8 -*-
"""
Buduje prezentacje Dobra Kaloria z opisu JSON na bazie szablonu DK_WZOR.pptx.

Uzycie:
    python build_deck.py spec.json [wyjscie.pptx]

Szablon (assets/DK_WZOR.pptx) ma osadzone fonty Mindset i Mindset Slim,
tlo, logo i grafiki pedzla (EMF). Nowe slajdy powstaja przez klonowanie
elementow z 5 slajdow szablonu, potem slajdy szablonu sa usuwane.

Jednostki w spec: centymetry (slajd 33.867 x 19.05 cm), rotacja w stopniach.
Opis formatu: references/spec-format.md
"""
import copy
import json
import os
import sys

from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
TEMPLATE = os.path.join(ROOT, "assets", "DK_WZOR.pptx")
FONT_FILES = {
    "Mindset": os.path.join(ROOT, "assets", "fonts", "Mindset.otf"),
    "Mindset Slim": os.path.join(ROOT, "assets", "fonts", "MindsetSlim.otf"),
    "Lato Bold": os.path.join(ROOT, "assets", "fonts", "Lato-Bold.ttf"),  # pomiar podpisu filmu (link_overlay)
    "Lato": os.path.join(ROOT, "assets", "fonts", "Lato-Regular.ttf"),  # długie teksty (29.09: Mindset Slim WERSALIKAMI nieczytelny)
}

# --- tokeny (z szablonu i 3 gotowych prezentacji) ---------------------------
SLIDE_W, SLIDE_H = 12192000, 6858000
CREAM = "FCF3E9"        # tlo slajdow tresci
GREEN = "006400"        # pasy, tlo koncowego, tekst punktow (wzor)
TEXT_GREEN = "005900"   # tekst w prezentacjach produktowych
DARK_GREEN = "084C1D"   # tytul okladki (wzor)
LIGHT_GREEN = "11953B"  # podtytul, strzalki, akcent
DOT_GREEN = "009114"    # kropki punktow (wzor)
SHADOW = "461600"       # cien pod packshotem (owal, softEdge 10 pt)
CM = 360000
PANEL_COVER_W = 6251944  # zielona polowa okladki w prezentacjach produktowych
SIDEBAR_W = 1202634      # zielony pas z logo na slajdach tresci
CONTENT_X0 = 1650000     # lewa krawedz tresci (za pasem)
CONTENT_X1 = 11900000
BODY_Y0 = 1450000        # pod paskiem tytulu
BODY_Y1 = 6650000

_fonts = {}


def cm(v):
    return int(round(float(v) * CM))


def font(name, size_pt):
    key = (name, size_pt)
    if key not in _fonts:
        _fonts[key] = ImageFont.truetype(FONT_FILES[name], max(1, int(size_pt * 10)))
    return _fonts[key]


def text_w_emu(text, size_pt, name="Mindset"):
    """Szerokosc tekstu w EMU. Mindset jest wersalikowy - dla niego mierzymy wersaliki."""
    t = text.upper() if name.startswith("Mindset") else text
    px = font(name, size_pt).getlength(t)  # rozmiar x10 -> px = 0.1 pt
    return int(px / 10 * 12700)


def fit_size(text, max_pt, min_pt, width_emu, name="Mindset"):
    s = max_pt
    while s > min_pt and text_w_emu(text, s, name) > width_emu:
        s -= 1
    return s


ZAWIESZKI = set("aiouwzAIOUWZ")  # jednoliterowe spójniki/przyimki - nigdy na końcu linii (typografia PL, 29.09)
DASHES = {"-", "–", "—"}


def typo_tokens(text):
    """Słowa do łamania z regułami polskiej typografii: 'w domu', 'i smak' = jeden token (zawieszka klejona do
    następnego słowa), myślnik klejony do poprzedniego (linia nie zaczyna się od '–')."""
    words = text.replace("\u00a0", " ").split(" ")
    words = [w for w in words if w]
    out = []
    for w in words:
        if out and (w in DASHES or out[-1] in ZAWIESZKI or out[-1].endswith(tuple(" " + z for z in ZAWIESZKI))):
            out[-1] = out[-1] + " " + w
        else:
            out.append(w)
    return out


def wrap_lines(text, size_pt, width_emu, name="Mindset"):
    """Łamanie na spacjach z typografią PL: bez zawieszek (a, i, o, u, w, z) na końcu linii, bez myślnika na
    początku linii i bez sieroty (jedno słowo w ostatniej linii) - reguła usera 29.09: zawsze."""
    words, lines, cur = typo_tokens(text), [], ""
    for w in words:
        t = (cur + " " + w).strip()
        if cur and text_w_emu(t, size_pt, name) > width_emu:
            lines.append(cur)
            cur = w
        else:
            cur = t
    if cur:
        lines.append(cur)
    if len(lines) >= 2 and len(lines[-1].split(" ")) == 1 and (
            len(lines[-2].split(" ")) >= 3 or len(lines[-1].strip("?!.,")) <= 4):  # sierota: krótkie słowo / długa linia
        prev = typo_tokens(lines[-2])
        moved = prev[-1] + " " + lines[-1]
        if len(prev) >= 2 and text_w_emu(moved, size_pt, name) <= width_emu:  # tylko gdy się zmieści
            lines[-2], lines[-1] = " ".join(prev[:-1]), moved
    return lines or [""]


def typo_nbsp(text, glue_last=True):
    """Tekst dla PowerPointa (fonty z twardą spacją, np. Lato): spacja po zawieszce -> twarda spacja, ostatnie dwa
    słowa akapitu sklejone (PowerPoint sam nie zostawi sieroty ani zawieszki, także po edycji). NIE dla Mindset
    (brak znaku U+00A0 w foncie) - nagłówki łamiemy sami w wrap_lines."""
    toks = typo_tokens(text)
    toks = [t.replace(" ", "\u00a0") for t in toks]
    if glue_last and len(toks) >= 3:
        toks[-2:] = [toks[-2] + "\u00a0" + toks[-1]]
    return " ".join(toks)


# --- klonowanie z szablonu --------------------------------------------------
class Deck:
    def __init__(self):
        self.prs = Presentation(TEMPLATE)
        self.tpl = list(self.prs.slides)  # 0 okladka, 1 punkty, 2 strzalki, 3 ozdobniki, 4 koniec
        self.n_tpl = len(self.tpl)

    def find(self, tpl_idx, name):
        def walk(shapes):
            for sh in shapes:  # najpierw ten poziom, potem grupy
                if sh.name == name:
                    return sh
            for sh in shapes:
                if sh.shape_type == 6:
                    r = walk(sh.shapes)
                    if r is not None:
                        return r
            return None
        r = walk(self.tpl[tpl_idx].shapes)
        if r is None:
            raise KeyError(name)
        return r

    def new_slide(self, bg=CREAM):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[0])
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string(bg)
        return s

    def clone(self, tpl_idx, name, slide):
        src = self.find(tpl_idx, name)
        src_part = self.tpl[tpl_idx].part
        el = copy.deepcopy(src._element)
        for node in el.iter():
            for attr in (qn("r:embed"), qn("r:link"), qn("r:id")):
                rid = node.get(attr)
                if rid:
                    part = src_part.related_part(rid)
                    node.set(attr, slide.part.relate_to(part, src_part.rels[rid].reltype))
        slide.shapes._spTree.insert_element_before(el, "p:extLst")
        nid = max([int(x.get("id")) for x in slide.shapes._spTree.iter(qn("p:cNvPr"))] + [1])
        for c in el.iter(qn("p:cNvPr")):
            nid += 1
            c.set("id", str(nid))
        return slide.shapes[-1]

    def finish(self, out):
        sld_ids = self.prs.slides._sldIdLst
        for sid in list(sld_ids)[: self.n_tpl]:
            self.prs.part.drop_rel(sid.get(qn("r:id")))
            sld_ids.remove(sid)
        import guard  # nie nadpisuj plików edytowanych ręcznie (lekcja 28.09.2026)
        out = guard.safe_target(out)
        self.prs.save(out)
        guard.record(out)
        return out


# --- tekst ------------------------------------------------------------------
def add_text(slide, x, y, w, h, paras, size, color=TEXT_GREEN, name="Mindset",
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, line_pt=None):
    """paras: lista napisow albo list [(tekst, font, kolor)] dla miksu w akapicie."""
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0  # szerokość pola = szerokość łamania (PowerPoint łamał inaczej niż my)
    if name and name.startswith("Mindset") and all(isinstance(p, str) for p in paras):
        # Mindset nie ma twardej spacji -> łamiemy sami z typografią PL (bez zawieszek i sierot, 29.09)
        avail = int(w * 0.97)
        paras = [ln for p in paras for ln in wrap_lines(p, size, avail, name)] if avail > 0 else paras
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        if line_pt:
            para.line_spacing = Pt(line_pt)
        runs = p if isinstance(p, list) else [(p, name, color)]
        for txt, fnt, col in runs:
            r = para.add_run()
            r.text = typo_nbsp(txt) if fnt and fnt.startswith("Lato") else txt
            r.font.size = Pt(size)
            r.font.name = fnt
            r.font.color.rgb = RGBColor.from_string(col)
    return tb


def set_text(shape, text, size=None):
    """Podmienia tekst w sklonowanym polu, zachowujac formatowanie 1. przebiegu."""
    txBody = shape._element.find(qn("p:txBody"))
    ps = txBody.findall(qn("a:p"))
    p0 = ps[0]
    for p in ps[1:]:
        txBody.remove(p)
    runs = p0.findall(qn("a:r"))
    for r in runs[1:]:
        p0.remove(r)
    r0 = runs[0]
    r0.find(qn("a:t")).text = text
    if size:
        r0.find(qn("a:rPr")).set("sz", str(int(size * 100)))
    end = p0.find(qn("a:endParaRPr"))
    if end is not None:
        p0.remove(end)


def bullets_box(slide, x, y, w, items, size=24, color=TEXT_GREEN):
    """Punkty jak w prezentacjach produktowych: kropka Arial 150%, interlinia 35 pt przy 24 pt."""
    line_pt = round(size * 35 / 24)
    n_lines = sum(len(wrap_lines(t, size, w - 342900 - 180000)) for t in items)
    h = int(n_lines * line_pt * 12700 + 180000)
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, t in enumerate(items):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pPr = para._p.get_or_add_pPr()
        pPr.set("marL", "342900")
        pPr.set("indent", "-342900")
        pPr.append(etree.fromstring(
            '<a:lnSpc xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            '<a:spcPts val="%d"/></a:lnSpc>' % (line_pt * 100)))
        pPr.append(etree.fromstring(
            '<a:buClr xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            '<a:srgbClr val="%s"/></a:buClr>' % color))
        pPr.append(etree.fromstring(
            '<a:buSzPct xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" val="150000"/>'))
        pPr.append(etree.fromstring(
            '<a:buFont xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="Arial"/>'))
        pPr.append(etree.fromstring(
            '<a:buChar xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" char="&#8226;"/>'))
        for j, ln in enumerate(wrap_lines(t, size, w - 342900 - 180000)):  # typografia PL: własne łamanie (29.09)
            if j:
                para.add_line_break()
            r = para.add_run()
            r.text = ln
            r.font.size = Pt(size)
            r.font.name = "Mindset"
            r.font.color.rgb = RGBColor.from_string(color)
    return tb, h


def fit_bullets(items, w, max_h, size=24, min_size=16):
    while size > min_size:
        line_pt = round(size * 35 / 24)
        n = sum(len(wrap_lines(t, size, w - 342900 - 180000)) for t in items)
        if n * line_pt * 12700 + 180000 <= max_h:
            break
        size -= 1
    return size


# --- obrazy -----------------------------------------------------------------
def alpha_bbox(path):
    """Ulamkowy bbox nieprzezroczystej czesci PNG (0..1)."""
    im = Image.open(path)
    if im.mode in ("RGBA", "LA") or "transparency" in im.info:
        bb = im.convert("RGBA").split()[-1].point(lambda a: 255 if a > 40 else 0).getbbox()
        if bb:
            W, H = im.size
            return bb[0] / W, bb[1] / H, bb[2] / W, bb[3] / H
    return 0.0, 0.0, 1.0, 1.0


def img_ratio(path):
    im = Image.open(path)
    return im.size[0] / im.size[1]


def add_shadow(slide, cx, cy, w, h, rot=0, color=SHADOW):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int(cx - w / 2)), Emu(int(cy - h / 2)), Emu(int(w)), Emu(int(h)))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(color)
    sh.line.fill.background()
    sh.shadow.inherit = False
    spPr = sh._element.spPr
    for old in spPr.findall(qn("a:effectLst")):
        spPr.remove(old)
    spPr.append(etree.fromstring(
        '<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:softEdge rad="127000"/></a:effectLst>'))
    sh.rotation = rot
    sh.name = "Cien"
    return sh


def place_image(slide, path, x, y, h=None, w=None, rot=0.0, shadow=False, shadow_color=SHADOW):
    r = img_ratio(path)
    if w is None:
        w = h * r
    if h is None:
        h = w / r
    if shadow:
        bl, bt, br, bb = alpha_bbox(path)
        vis_w = (br - bl) * w
        sw, shh = vis_w * 0.95, max(vis_w * 0.14, 250000)
        add_shadow(slide, x + (bl + br) / 2 * w, y + bb * h - shh * 0.2, sw, shh, rot=rot, color=shadow_color)
    pic = slide.shapes.add_picture(path, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    pic.rotation = rot
    return pic


def trio_row(slide, items, region):
    """3 paczki jak w nowym stylu: dwie z tyłu po bokach, środkowa z przodu, lekko nachodzą (29.09 - trójkąt
    z auto_packshots był drobny i rozrzucony). items: [lewa, środek(przód), prawa]. Zwraca bbox."""
    x0, y0, x1, y1 = region
    rw, rh = x1 - x0, y1 - y0
    hc = rh
    r = [img_ratio(it["path"]) for it in items]
    while True:
        hs = hc * 0.84
        wl, wc, wr = hs * r[0], hc * r[1], hs * r[2]
        tot = wl * 0.62 + wc + wr * 0.62
        if tot <= rw:
            break
        hc *= 0.95
    cx = x0 + rw / 2
    xc = cx - wc / 2
    yc = y0 + (rh - hc) / 2
    ys = yc + hc * 0.04
    slots = [(items[0], xc - wl * 0.62, ys, hs, -5), (items[2], xc + wc - wr * 0.38, ys, hs, 5),
             (items[1], xc, yc, hc, 0)]
    for it, x, y, h, rot in slots:
        pic = place_image(slide, it["path"], x, y, h=h, rot=rot, shadow=True,
                          shadow_color=it.get("shadow_color", SHADOW))
        if it.get("name"):
            pic.name = it["name"]
    return (xc - wl * 0.62, yc, xc + wc + wr * 0.62, yc + hc)


def auto_packshots(slide, items, region):
    """Uklada packshoty w obszarze (x0,y0,x1,y1) wg wzorcow z prezentacji produktowych."""
    x0, y0, x1, y1 = region
    rw, rh = x1 - x0, y1 - y0
    n = len(items)
    slots = []
    if n == 1:
        h = rh * 0.88
        slots = [(x0 + rw / 2, y0 + rh * 0.5, h, 0)]
    elif n == 2:  # kaskada: tyl prawo-gora, przod lewo-dol (Cynamonka, Kulki)
        h = rh * 0.80
        slots = [(x0 + rw * 0.66, y0 + rh * 0.42, h, 4), (x0 + rw * 0.36, y0 + rh * 0.58, h, -6)]
    elif n == 3:  # trojkat: dwa z tylu, jeden z przodu (Kremy)
        h = rh * 0.58
        slots = [(x0 + rw * 0.22, y0 + rh * 0.40, h, 0), (x0 + rw * 0.78, y0 + rh * 0.40, h, 0),
                 (x0 + rw * 0.50, y0 + rh * 0.62, h, 0)]
    else:
        h = rh * 0.55
        step = rw / n
        slots = [(x0 + step * (i + 0.5), y0 + rh * (0.45 if i % 2 else 0.55), h, 0) for i in range(n)]
    placed = []
    for it, (cx, cy, h, rot) in zip(items, slots):
        h = cm(it["h"]) if "h" in it else h * it.get("scale", 1.0)
        w = h * img_ratio(it["path"])
        if "x" in it:
            x, y = cm(it["x"]), cm(it["y"])
        else:
            x, y = cx - w / 2, cy - h / 2
        rot = it.get("rot", rot)
        pic = place_image(slide, it["path"], x, y, h=h, rot=rot, shadow=it.get("shadow", True),
                          shadow_color=it.get("shadow_color", SHADOW))
        if it.get("name"):
            pic.name = it["name"]
        placed.append((x, y, x + w, y + h))
    return placed


def add_props(slide, props, bbox):
    """Rekwizyty (owoce, kawalki produktu) - jawnie x/y albo 'anchor' wzgledem packshotow."""
    if not props:
        return
    bx0, by0, bx1, by1 = bbox
    anchors = {"tl": (bx0, by0), "tr": (bx1, by0), "bl": (bx0, by1), "br": (bx1, by1),
               "l": (bx0, (by0 + by1) / 2), "r": (bx1, (by0 + by1) / 2), "b": ((bx0 + bx1) / 2, by1)}
    for p in props:
        w = cm(p.get("w", 3))
        h = w / img_ratio(p["path"])
        if "x" in p:
            x, y = cm(p["x"]), cm(p["y"])
        else:
            ax, ay = anchors[p.get("anchor", "br")]
            x, y = ax - w / 2 + cm(p.get("dx", 0)), ay - h / 2 + cm(p.get("dy", 0))
        x = min(max(x, SIDEBAR_W + 50000), SLIDE_W - w * 0.3)
        place_image(slide, p["path"], x, y, w=w, rot=p.get("rot", 0), shadow=p.get("shadow", False))


# --- elementy wspolne -------------------------------------------------------
def add_morph(slide, dur=1200):
    """Przejscie Morph (PowerPoint 2019+/365), fallback Fade. Obiekty o tych samych nazwach '!!...'
    na sasiednich slajdach plynnie sie przesuwaja/skaluja; reszta sie przenika."""
    pns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    el = etree.fromstring(
        '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:p="%s" xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" '
        'xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main">'
        '<mc:Choice Requires="p159"><p:transition spd="slow" p14:dur="%d"><p159:morph option="byObject"/>'
        '</p:transition></mc:Choice><mc:Fallback><p:transition spd="slow"><p:fade/></p:transition>'
        '</mc:Fallback></mc:AlternateContent>' % (pns, dur))
    anchor = slide._element.find("{%s}clrMapOvr" % pns)
    if anchor is None:
        anchor = slide._element.find("{%s}cSld" % pns)
    anchor.addnext(el)


def content_frame(deck, slide, title):
    deck.clone(1, "Prostokąt 33", slide).name = "!!pas"          # zielony pas (stoi w miejscu przy Morph)
    deck.clone(1, "Obraz 7", slide).name = "!!logo_pas"          # logo w pasie
    grp = deck.clone(1, "Grupa 3", slide)      # pedzel + tytul
    grp.name = "!!tytul"
    tbox = [s for s in grp.shapes if s.has_text_frame][0]
    size = fit_size(title, 44, 26, int(tbox.width * 0.92))
    set_text(tbox, title, size)


def add_callouts(deck, slide, callouts):
    """Ozdobniki z slajdu 4 szablonu: speech (dymek), banner (strzalka), round (baner zaokraglony),
    box (ramka), bubble (maly dymek), polska (mapka + 'Polska firma rodzinna')."""
    names = {"banner": "Obraz 2", "round": "Obraz 13", "box": "Obraz 5", "bubble": "Obraz 11",
             "speech": "Obraz 15", "polska": "Grupa 12"}
    for c in callouts or []:
        kind = c["kind"]
        if kind == "polska":  # w szablonie grupa lezy na zielonej ramce 'box' - bez niej bialy tekst znika
            gw = cm(c.get("w", 2351726 / CM))
            gh = int(gw * 707886 / 2351726)
            bx = deck.clone(3, "Obraz 5", slide)
            bx.name = "tlo-ramka"
            bx.left, bx.top = cm(c["x"]) - int(gw * 0.118), cm(c["y"]) - int(gh * 0.46)
            bx.width, bx.height = int(gw * 1.178), int(gh * 1.80)
        sh = deck.clone(3, names[kind], slide)
        ratio = sh.width / sh.height
        w = cm(c.get("w", sh.width / CM))
        h = cm(c["h"]) if "h" in c else int(w / ratio)
        sh.left, sh.top, sh.width, sh.height = cm(c["x"]), cm(c["y"]), w, h
        if kind == "polska":
            if "text" in c:
                set_text([s for s in sh.shapes if s.has_text_frame][0], c["text"])
            continue
        if c.get("text"):
            pad = int(w * 0.10)
            body_h = h * (0.82 if kind in ("speech", "bubble") else 1.0)
            paras = c["text"] if isinstance(c["text"], list) else [c["text"]]
            size = c.get("size", 24)
            add_text(slide, sh.left + pad, sh.top, w - 2 * pad, int(body_h), paras, size,
                     color="FFFFFF", name=c.get("font", "Mindset"), anchor=MSO_ANCHOR.MIDDLE)


# --- typy slajdow -----------------------------------------------------------
def s_cover(deck, spec):
    variant = spec.get("variant", "product")
    s = deck.new_slide()
    if variant == "template":  # dokladnie jak wzor: TYTUL / PREZENTACJI / linia / podtytul
        deck.clone(0, "Prostokąt 9", s)
        deck.clone(0, "Obraz 4", s)
        g = deck.clone(0, "Grupa 1", s)
        tb = {sh.name: sh for sh in g.shapes[0].shapes}
        set_text(tb["pole tekstowe 14"], spec.get("title", "Tytuł"), fit_size(spec.get("title", ""), 115, 60, 5900000))
        set_text(tb["pole tekstowe 15"], spec.get("title2", ""), fit_size(spec.get("title2", ""), 48, 28, 5900000))
        sub = [x for x in g.shapes if x.has_text_frame][0]
        set_text(sub, spec.get("subtitle", ""), fit_size(spec.get("subtitle", ""), 44, 24, 5900000))
        return
    # wariant produktowy / raportowy: szerszy panel, stos napisow w prawej polowie
    panel = deck.clone(0, "Prostokąt 9", s)
    panel.width = PANEL_COVER_W
    logo = deck.clone(0, "Obraz 4", s)
    logo.left = int((PANEL_COVER_W - logo.width) / 2)
    x0, x1 = PANEL_COVER_W, SLIDE_W
    width = x1 - x0 - 300000
    blocks = []  # (tekst, rozmiar, kolor)
    if spec.get("kicker"):
        blocks.append((spec["kicker"], fit_size(spec["kicker"], 75, 40, width), TEXT_GREEN))
    if spec.get("title"):
        blocks.append((spec["title"], fit_size(spec["title"], 60, 30, width), TEXT_GREEN))
    subs = spec.get("subtitle", [])
    subs = subs if isinstance(subs, list) else [subs]
    for i, t in enumerate(subs):
        max_pt = 50 if len(subs) == 1 and len(t) <= 22 else 28
        col = LIGHT_GREEN if spec.get("alternate") and i % 2 else TEXT_GREEN
        blocks.append((t, fit_size(t, max_pt, 18, width), col))
    heights = [int(sz * 1.18 * 12700) for _, sz, _ in blocks]
    line_h = 293331
    gap = 150000
    total = sum(heights) + gap + line_h
    y = int(SLIDE_H * 0.47 - total / 2)
    if spec.get("images"):  # 29.09: okładka z produktem - napisy wyżej, pod nimi cała linia (3 paczki)
        y = int(SLIDE_H * 0.07)
    for (t, sz, col), hh in zip(blocks, heights):
        tb = add_text(s, x0 + 150000, y, width, hh, [t], sz, color=col)
        tb.text_frame.margin_top = tb.text_frame.margin_bottom = 0
        y += hh
    line = deck.clone(0, "Obraz 2", s)   # pociagniecie pedzla pod tytulem
    line.width, line.height = 3205935, line_h
    line.left = int(x0 + (x1 - x0 - line.width) / 2)
    line.top = y + gap
    if spec.get("images"):
        top = line.top + line_h + cm(0.5)
        reg = (x0 + cm(1.2), top, x1 - cm(1.2), SLIDE_H - cm(0.5))
        (trio_row if len(spec["images"]) == 3 else auto_packshots)(s, spec["images"], reg)


def s_product(deck, spec):
    """Tytul w pedzlu + punkty po lewej + packshoty z cieniami i rekwizytami po prawej."""
    s = deck.new_slide()
    content_frame(deck, s, spec["title"])
    images = spec.get("images", [])
    bullets = spec.get("bullets", [])
    if not images:  # punkty w stylu wzoru: zielone kropki, 32 pt, cala szerokosc
        return s_bullets_full(deck, s, bullets)
    text_x = cm(spec.get("text_x", 3.9))
    text_w = cm(spec.get("text_w", 17.5))
    size = spec.get("size") or fit_bullets(bullets, text_w, BODY_Y1 - BODY_Y0, 24)
    _, h = bullets_box(s, text_x, 0, text_w, bullets, size)
    s.shapes[-1].top = int(BODY_Y0 + (BODY_Y1 - BODY_Y0 - h) / 2 + spec.get("text_dy", 0) * CM)
    region = spec.get("region")
    region = tuple(cm(v) for v in region) if region else (cm(18.2), BODY_Y0 - 50000, SLIDE_W - 150000, SLIDE_H - 100000)
    behind = [p for p in spec.get("props", []) if p.get("behind")]
    front = [p for p in spec.get("props", []) if not p.get("behind")]
    rx0, ry0, rx1, ry1 = region
    add_props(s, behind, region)
    if len(images) == 3:  # cała linia: rząd 3 paczek (29.09)
        bbox = trio_row(s, images, (rx0 + cm(0.3), ry0 + cm(0.9), rx1 - cm(0.6), ry1 - cm(0.9)))
    else:
        placed = auto_packshots(s, images, region)
        bbox = (min(p[0] for p in placed), min(p[1] for p in placed), max(p[2] for p in placed),
                max(p[3] for p in placed))
    add_props(s, front, bbox)
    add_callouts(deck, s, spec.get("callouts"))


def s_bullets_full(deck, s, bullets):
    y, step = 1700000, 870000
    size = 32
    while size > 20 and max(text_w_emu(b, size) for b in bullets) > 9453465:
        size -= 1
    step = int(870000 * size / 32) if len(bullets) <= 5 else int((BODY_Y1 - 1700000) / len(bullets))
    for b in bullets:
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(1850235), Emu(y + 178000), Emu(228600), Emu(228600))
        dot.fill.solid()
        dot.fill.fore_color.rgb = RGBColor.from_string(DOT_GREEN)
        dot.line.fill.background()
        dot.shadow.inherit = False
        add_text(s, 2173406, y, 9453465, 584775, [b], size, color=GREEN, align=PP_ALIGN.LEFT)
        y += step


def s_bullets(deck, spec):
    s = deck.new_slide()
    content_frame(deck, s, spec["title"])
    s_bullets_full(deck, s, spec.get("bullets", []))
    add_callouts(deck, s, spec.get("callouts"))


def s_arrows(deck, spec):
    """Strzalki 11953B + tekst 36 pt; element {'text':..., 'slim': True} -> Mindset Slim."""
    s = deck.new_slide()
    content_frame(deck, s, spec["title"])
    n_items = len(spec.get("items", []))
    y = int(max(1850000, BODY_Y0 + (BODY_Y1 - BODY_Y0 - n_items * 832159) / 2))
    for it in spec.get("items", []):
        it = it if isinstance(it, dict) else {"text": it}
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Emu(2278117), Emu(y), Emu(770021), Emu(505327))
        ar.fill.solid()
        ar.fill.fore_color.rgb = RGBColor.from_string(LIGHT_GREEN)
        ar.line.fill.background()
        ar.shadow.inherit = False
        name = "Mindset Slim" if it.get("slim") else "Mindset"
        size = fit_size(it["text"], 36, 20, 8600000, name)
        add_text(s, 3180000, y - 30000, 8600000, 553549, [it["text"]], size, color=GREEN, name=name, align=PP_ALIGN.LEFT)
        y += 832159
    add_callouts(deck, s, spec.get("callouts"))


def s_gallery(deck, spec):
    """Siatka obrazow (zrzuty z telefonu, materialy tworcow) + opcjonalny opis obok (styl raportu DATESY)."""
    s = deck.new_slide()
    content_frame(deck, s, spec["title"])
    imgs = spec.get("images", [])
    caption = spec.get("caption")
    x0, x1 = CONTENT_X0, CONTENT_X1 - (cm(spec.get("caption_w", 7)) if caption else 0)
    y0, y1 = BODY_Y0 + 150000, BODY_Y1
    n = len(imgs)
    cols = spec.get("cols") or min(n, 4)
    rows = (n + cols - 1) // cols
    gap = cm(spec.get("gap", 0.5))
    cw, ch = (x1 - x0 - gap * (cols - 1)) / cols, (y1 - y0 - gap * (rows - 1)) / rows
    for i, it in enumerate(imgs):
        it = it if isinstance(it, dict) else {"path": it}
        r = img_ratio(it["path"])
        w, h = (cw, cw / r) if cw / r <= ch else (ch * r, ch)
        cx = x0 + (i % cols) * (cw + gap) + (cw - w) / 2
        cy = y0 + (i // cols) * (ch + gap) + (ch - h) / 2
        if spec.get("frame", False):  # czarna ramka "telefonu"
            fr = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(cx - 60000)), Emu(int(cy - 60000)),
                                    Emu(int(w + 120000)), Emu(int(h + 120000)))
            fr.adjustments[0] = 0.08
            fr.fill.solid()
            fr.fill.fore_color.rgb = RGBColor.from_string("1A1A1A")
            fr.line.fill.background()
            fr.shadow.inherit = False
        s.shapes.add_picture(it["path"], Emu(int(cx)), Emu(int(cy)), Emu(int(w)), Emu(int(h)))
    if caption:
        paras = caption if isinstance(caption, list) else [caption]
        add_text(s, x1 + 250000, y0, CONTENT_X1 - x1 - 250000, y1 - y0, paras, spec.get("caption_size", 14),
                 color=TEXT_GREEN, name="Mindset Slim", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
    add_callouts(deck, s, spec.get("callouts"))


def s_free(deck, spec):
    """Pusty slajd z ramka (pas + tytul): tylko obrazy/ozdobniki/teksty podane jawnie."""
    s = deck.new_slide()
    content_frame(deck, s, spec["title"])
    for it in spec.get("images", []):
        place_image(s, it["path"], cm(it["x"]), cm(it["y"]), h=cm(it["h"]) if "h" in it else None,
                    w=cm(it["w"]) if "w" in it else None, rot=it.get("rot", 0), shadow=it.get("shadow", False))
    for t in spec.get("texts", []):
        add_text(s, cm(t["x"]), cm(t["y"]), cm(t["w"]), cm(t.get("h", 2)), t["text"] if isinstance(t["text"], list) else [t["text"]],
                 t.get("size", 24), color=t.get("color", TEXT_GREEN), name=t.get("font", "Mindset"),
                 align={"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[t.get("align", "left")])
    add_callouts(deck, s, spec.get("callouts"))


def s_end(deck, spec):
    s = deck.new_slide(bg=GREEN)
    deck.clone(4, "Obraz 5", s)
    tb = deck.clone(4, "pole tekstowe 1", s)
    tag = spec.get("hashtag")
    if tag:  # domyslnie #zawszedobra zostaje z szablonu
        runs = tb.text_frame.paragraphs[0].runs
        runs[1].text = tag.lstrip("#")


def add_source(slide, text):
    """Stopka ze zrodlem danych (obowiazkowa na slajdach z liczbami)."""
    if text:
        add_text(slide, CONTENT_X0, SLIDE_H - 420000, CONTENT_X1 - CONTENT_X0, 300000, [text], 14,
                 color=TEXT_GREEN, name="Mindset Slim", align=PP_ALIGN.LEFT)


def s_stats(deck, spec):
    """Duze liczby w zielonych ramkach-pedzlach (ozdobnik 'box' ze wzoru) + opis pod spodem.
    items: [{value, label, note?}] (2-4). Opcjonalnie images -> packshot po prawej."""
    s = deck.new_slide()
    content_frame(deck, s, spec["title"])
    items = spec["items"]
    has_img = bool(spec.get("images"))
    x0, x1 = CONTENT_X0 + 100000, (cm(21.5) if has_img else CONTENT_X1)
    if spec.get("lead"):
        add_text(s, x0, BODY_Y0 + 150000, x1 - x0, 500000, [spec["lead"]], 24, color=TEXT_GREEN,
                 name="Mindset Slim", align=PP_ALIGN.CENTER)
    n = len(items)
    gap = 250000
    cw = (x1 - x0 - gap * (n - 1)) / n
    box_w = min(cw, 3300000)
    box_h = int(box_w / 2.177)
    lead_h = 700000 if spec.get("lead") else 0
    block = box_h + 200000 + 800000 + (450000 if any(it.get("note") for it in items) else 0)
    top = BODY_Y0 + lead_h + max(250000, (BODY_Y1 - 350000 - BODY_Y0 - lead_h - block) / 2)
    if spec.get("lead"):
        s.shapes[-1].top = int(top - lead_h)
    for i, it in enumerate(items):
        cx = x0 + i * (cw + gap) + cw / 2
        box = deck.clone(3, "Obraz 5", s)
        box.name = "tlo-ramka"
        box.left, box.top, box.width, box.height = int(cx - box_w / 2), int(top), int(box_w), box_h
        vsize = fit_size(it["value"], 72, 36, int(box_w * 0.8))
        tb = add_text(s, int(cx - box_w / 2), int(top), int(box_w), box_h, [it["value"]], vsize,
                      color="FFFFFF", anchor=MSO_ANCHOR.MIDDLE)
        ly = int(top + box_h + 200000)
        lz = 22
        while lz > 16 and len(wrap_lines(it["label"], lz, int(cw * 0.95))) > 2:  # etykieta max 2 linie (29.09)
            lz -= 1
        lines = wrap_lines(it["label"], lz, int(cw * 0.95))
        lh = int(len(lines) * lz * 1.25 * 12700) + 60000
        add_text(s, int(cx - cw / 2), ly, int(cw), lh, lines, lz, color=TEXT_GREEN)
        if it.get("note"):
            add_text(s, int(cx - cw / 2), ly + lh + 60000, int(cw), 900000, [it["note"]], 14,
                     color=LIGHT_GREEN, name="Mindset Slim")
    if has_img:
        auto_packshots(s, spec["images"], (cm(22.5), BODY_Y0, SLIDE_W - 200000, SLIDE_H - 250000))
    add_source(s, spec.get("source"))
    add_callouts(deck, s, spec.get("callouts"))


def s_compare(deck, spec):
    """Warianty obok siebie (np. opakowania z badania) z wynikiem pod kazdym.
    items: [{path, label, value, sub?}], highlight: indeks zwyciezcy."""
    s = deck.new_slide()
    content_frame(deck, s, spec["title"])
    items = spec["items"]
    hi = spec.get("highlight")
    n = len(items)
    x0, x1 = CONTENT_X0 + 100000, CONTENT_X1
    cw = (x1 - x0) / n
    img_top, img_h = BODY_Y0 + 250000, cm(spec.get("img_h", 7.2))
    for i, it in enumerate(items):
        cx = x0 + cw * (i + 0.5)
        scale = 1.08 if i == hi else 0.92
        h = img_h * scale
        w = h * img_ratio(it["path"])
        y = img_top + (img_h - h)
        place_image(s, it["path"], cx - w / 2, y, h=h, shadow=True)
        vy = img_top + img_h + 300000
        if not it.get("value"):
            pass
        elif i == hi:
            box = deck.clone(3, "Obraz 5", s)
            box.name = "tlo-ramka"
            bw = min(cw * 0.9, 2600000)
            box.left, box.top, box.width, box.height = int(cx - bw / 2), int(vy - 80000), int(bw), int(bw / 2.177 * 0.62)
            box.height = 820000
            add_text(s, int(cx - bw / 2), int(vy - 80000), int(bw), 820000, [it["value"]], 48, color="FFFFFF",
                     anchor=MSO_ANCHOR.MIDDLE)
        else:
            add_text(s, int(cx - cw / 2), int(vy - 80000), int(cw), 820000, [it["value"]], 44, color=TEXT_GREEN,
                     anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, int(cx - cw / 2), int(vy + 820000), int(cw), 400000, [it["label"]], 20,
                 color=TEXT_GREEN if i == hi else LIGHT_GREEN)
        if it.get("sub"):
            add_text(s, int(cx - cw / 2), int(vy + 1180000), int(cw), 400000, [it["sub"]], 15,
                     color=TEXT_GREEN, name="Mindset Slim")
    add_source(s, spec.get("source"))
    add_callouts(deck, s, spec.get("callouts"))


def s_bars(deck, spec):
    """Poziomy wykres slupkowy: etykieta po lewej, wartosc na koncu slupka (bez osi i legendy).
    items: [{label, value(0-100)}]; highlight: indeksy wyrozniane ciemna zielenia."""
    s = deck.new_slide()
    content_frame(deck, s, spec["title"])
    items = spec["items"]
    hi = set(spec.get("highlight", []))
    has_img = bool(spec.get("images"))
    x0, x1 = CONTENT_X0 + 100000, (cm(22) if has_img else CONTENT_X1)
    label_w = cm(spec.get("label_w", 6.8))
    top = BODY_Y0 + (850000 if spec.get("lead") else 450000)
    if spec.get("lead"):
        add_text(s, x0, BODY_Y0 + 200000, x1 - x0, 500000, [spec["lead"]], 22, color=TEXT_GREEN,
                 name="Mindset Slim", align=PP_ALIGN.LEFT)
    avail = BODY_Y1 - 350000 - top
    row = min(avail / len(items), 820000)
    bar_h = row * 0.56
    bar_x = x0 + label_w
    bar_max = x1 - bar_x - 1000000
    vmax = spec.get("max", 100)
    for i, it in enumerate(items):
        y = top + i * row
        add_text(s, x0, int(y), label_w - 150000, int(row), [it["label"]], 20, color=TEXT_GREEN,
                 align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        w = max(bar_max * it["value"] / vmax, 60000)
        col = GREEN if i in hi else "9CC08B"
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(bar_x)), Emu(int(y + (row - bar_h) / 2)),
                               Emu(int(w)), Emu(int(bar_h)))
        b.adjustments[0] = 0.5
        b.fill.solid()
        b.fill.fore_color.rgb = RGBColor.from_string(col)
        b.line.fill.background()
        b.shadow.inherit = False
        add_text(s, int(bar_x + w + 120000), int(y), 1300000, int(row), [it.get("text", "%d%%" % round(it["value"]))],
                 26 if i in hi else 22, color=TEXT_GREEN, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
    if has_img:
        auto_packshots(s, spec["images"], (cm(23), BODY_Y0, SLIDE_W - 200000, SLIDE_H - 250000))
    add_source(s, spec.get("source"))
    add_callouts(deck, s, spec.get("callouts"))


def s_skus(deck, spec):
    """Smaki / SKU w kolumnach: packshot z cieniem (albo kolo w kolorze smaku, gdy brak zdjecia),
    nazwa w zielonej ramce-pedzlu, pod spodem gramatura i EAN (Mindset Slim).
    items: [{name, meta, ean, color, path?}]"""
    s = deck.new_slide()
    content_frame(deck, s, spec["title"])
    items = spec["items"]
    n = len(items)
    x0, x1 = CONTENT_X0 + 100000, CONTENT_X1
    cw = (x1 - x0) / n
    img_cy, img_h = BODY_Y0 + cm(3.6), cm(6.4)
    for i, it in enumerate(items):
        cx = x0 + cw * (i + 0.5)
        if it.get("path"):
            h = img_h
            w = h * img_ratio(it["path"])
            place_image(s, it["path"], cx - w / 2, img_cy - h / 2, h=h, shadow=True,
                        shadow_color=it.get("shadow_color", SHADOW))
        else:
            d = cm(5.2)
            c = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int(cx - d / 2)), Emu(int(img_cy - d / 2)), Emu(int(d)), Emu(int(d)))
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor.from_string(it.get("color", "C9DABB"))
            c.line.fill.background()
            c.shadow.inherit = False
        by = img_cy + img_h / 2 + cm(0.5)
        bw = min(cw * 0.9, cm(8.6))
        bh = cm(1.9)
        box = deck.clone(3, "Obraz 5", s)
        box.name = "tlo-ramka"
        box.left, box.top, box.width, box.height = int(cx - bw / 2), int(by), int(bw), int(bh)
        size = fit_size(it["name"], 32, 18, int(bw * 0.84))
        add_text(s, int(cx - bw / 2), int(by), int(bw), int(bh), [it["name"]], size, color="FFFFFF",
                 anchor=MSO_ANCHOR.MIDDLE)
        info = [x for x in (it.get("meta"), ("EAN " + it["ean"]) if it.get("ean") else None) if x]
        add_text(s, int(cx - cw / 2), int(by + bh + cm(0.3)), int(cw), cm(2), info, 20, color=TEXT_GREEN,
                 name="Mindset Slim")
    add_source(s, spec.get("source"))
    add_callouts(deck, s, spec.get("callouts"))


def _round_geom(shape, radius_emu):
    """Zaokraglone rogi wlasna geometria (nie maska): prstGeom roundRect na zdjeciu/filmie."""
    adj = int(50000 * radius_emu / max(1, min(shape.width, shape.height)))
    g = shape._element.spPr.find(qn("a:prstGeom"))
    g.set("prst", "roundRect")
    for c in list(g):
        g.remove(c)
    av = etree.SubElement(g, qn("a:avLst"))
    gd = etree.SubElement(av, qn("a:gd"))
    gd.set("name", "adj")
    gd.set("fmla", "val %d" % adj)


def _card(slide, x, y, w, h, fill, name, radius=None):
    radius = cm(0.45) if radius is None else radius
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    c.adjustments[0] = radius / min(w, h)
    c.fill.solid()
    c.fill.fore_color.rgb = RGBColor.from_string(fill)
    c.line.fill.background()
    c.shadow.inherit = False
    c.name = name
    return c


def s_media(deck, spec):
    """Po lewej karta z twierdzeniem (male zdjecie + naglowek + akapit), po prawej karta z filmem
    (tytul + plakat 16:9 z zaokraglonymi rogami i przyciskiem play z linkiem).
    Pola: title, claim_image, claim_title, claim_text (albo text), media_title, poster, link, video?"""
    spec = dict(spec, media_title=spec.get("media_title", "").replace("\n", " "))  # tytuł filmu: styl A łamie sam
    s = deck.new_slide()
    content_frame(deck, s, spec["title"])
    x0, x1 = CONTENT_X0 + 100000, CONTENT_X1
    y0, y1 = BODY_Y0 + 250000, BODY_Y1 - 250000
    gap = cm(0.7)
    lw = int((x1 - x0 - gap) * 0.42)
    rx, rw = x0 + lw + gap, x1 - (x0 + lw + gap)
    h = y1 - y0
    pad = cm(0.7)
    # lewa karta
    _card(s, x0, y0, lw, h, "FFFFFF", "Karta - twierdzenie")
    iw = lw - 2 * pad
    ih = int(h * 0.24)  # 29.09: mniejsza kulka, więcej miejsca na czytelny opis (min. 15 pt)
    cy = y0 + pad
    if spec.get("claim_image"):
        r = img_ratio(spec["claim_image"])
        pw, ph = (ih * r, ih) if ih * r <= iw else (iw, iw / r)
        place_image(s, spec["claim_image"], x0 + (lw - pw) / 2, cy + (ih - ph) / 2, w=pw)
    ty = cy + ih + cm(0.3)
    tsize = fit_size(spec["claim_title"], 26, 20, iw * 2 - cm(1))
    parts = spec["claim_title"].split("\n")  # ręczne łamanie ze specu ("1 g kreatyny\nw każdej kulce")
    tl = sum(len(wrap_lines(pt, tsize, iw)) for pt in parts)
    th = int(tl * tsize * 1.25 * 12700) + 60000
    add_text(s, x0 + pad, ty, iw, th, parts, tsize, color=TEXT_GREEN, name="Mindset")
    by = ty + th + cm(0.3)
    bsize = 18
    while bsize > 15 and len(wrap_lines(spec.get("claim_text", spec.get("text", "")), bsize, iw - 60000, "Lato")) * bsize * 1.3 * 12700 > y1 - pad - by:
        bsize -= 1
    add_text(s, x0 + pad, by, iw, int(y1 - pad - by), [spec.get("claim_text", spec.get("text", ""))], bsize, color=TEXT_GREEN,
             name="Lato", line_pt=round(bsize * 1.3))
    # prawa karta
    _card(s, rx, y0, rw, h, "F5E6D3", "Karta - film")
    mw = rw - 2 * pad
    msize = 20
    while msize > 16 and len(wrap_lines(spec["media_title"], msize, mw)) > 2:
        msize -= 1
    ml = min(2, len(wrap_lines(spec["media_title"], msize, mw)))
    mh = int(ml * msize * 1.25 * 12700) + 60000
    pw = mw
    ph = int(pw * 9 / 16)
    if mh + cm(0.4) + ph > h - 2 * pad:  # za wysoko - zmniejsz, zachowaj 16:9 i wysrodkuj
        ph = int(h - 2 * pad - mh - cm(0.4))
        pw = int(ph * 16 / 9)
    ty0 = y0 + (h - (mh + cm(0.4) + ph)) / 2  # blok tytul+film wysrodkowany w karcie
    add_text(s, rx + pad, int(ty0), mw, mh, [spec["media_title"]], msize, color=TEXT_GREEN, name="Mindset")
    py = ty0 + mh + cm(0.4)
    px = rx + (rw - pw) / 2
    link = spec.get("link")
    poster = spec.get("poster") or (yt_thumb(link, os.environ.get("DK_CACHE_DIR") or os.path.join(ROOT, "assets", "placeholdery"))
                                    if link else None)
    if not poster:  # brak miniatury: ciemna klatka zastepcza 16:9 (wtedy rysujemy przycisk play)
        import io
        buf = io.BytesIO()
        Image.new("RGB", (1280, 720), (24, 24, 24)).save(buf, "PNG")
        buf.seek(0)
        poster = buf
    if spec.get("video"):
        pic = s.shapes.add_movie(spec["video"], Emu(int(px)), Emu(int(py)), Emu(pw), Emu(ph),
                                 poster_frame_image=poster, mime_type="video/mp4")
    else:
        pic = s.shapes.add_picture(poster, Emu(int(px)), Emu(int(py)), Emu(pw), Emu(ph))
    _round_geom(pic, cm(0.45))
    pic.name = "Film - podmiana: wstaw nowy film, Malarz formatów z tego filmu na nowy"
    if link and not spec.get("video"):  # film otwiera się po kliknięciu na YouTube (błąd 153 w PowerPoint, 29.09)
        pic.name = "Film YouTube (link) - prawy klik: Edytuj link / Zmień obraz"
        pic.click_action.hyperlink.address = link
        link_overlay(s, px, py, pw, ph, link)
    add_source(s, spec.get("source"))


def s_article(deck, spec):
    """Tytul + jedna duza kremowa karta z dlugim tekstem (paragraphs), opcjonalnie props po bokach
    (2 z lewej, 2 z prawej, czesciowo nachodza na krawedz karty). props: sciezki PNG albo {path}."""
    s = deck.new_slide()
    content_frame(deck, s, spec["title"])
    props = [p["path"] if isinstance(p, dict) else p for p in spec.get("props", [])]
    inset = cm(1.3) if props else 0
    x0, x1 = CONTENT_X0 + 100000 + inset, CONTENT_X1 - inset
    y0, y1 = BODY_Y0 + 250000, BODY_Y1 - 250000
    padx = cm(2.6) if props else cm(1.2)
    pady = cm(1.0)
    tw = x1 - x0 - 2 * padx
    paras = spec["paragraphs"]
    gap_pt = 12
    size = 20  # 29.09: długi akapit w Lato (nie Mindset Slim wersalikami), czytelny z projektora
    while True:
        line_pt = round(size * 1.35)
        n = sum(len(wrap_lines(p, size, tw - 80000, "Lato")) for p in paras)
        need = n * line_pt * 12700 + gap_pt * (len(paras) - 1) * 12700
        if need <= y1 - y0 - 2 * pady or size <= 15:
            break
        size -= 1
    # karta dopasowana do tekstu (min. 10 cm), wysrodkowana w polu tresci
    ch = min(y1 - y0, max(int(need + 2 * pady), cm(10)))
    y0 = int((y0 + y1) / 2 - ch / 2)
    y1 = y0 + ch
    _card(s, x0, y0, x1 - x0, y1 - y0, "F5E6D3", "Karta - tekst", radius=cm(0.6))
    box = add_text(s, x0 + padx, y0 + pady, tw, y1 - y0 - 2 * pady, paras, size, color=TEXT_GREEN,
                   name="Lato", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, line_pt=line_pt)
    for para in box.text_frame.paragraphs[:-1]:
        para.space_after = Pt(gap_pt)
    if props:
        sizes = [cm(3.6), cm(2.8), cm(3.2), cm(3.6)]
        ymid = (y0 + y1) / 2
        slots = [(x0, ymid - cm(3.4), -1), (x0, ymid + cm(1.0), -1),
                 (x1, ymid - cm(3.6), 1), (x1, ymid + cm(0.8), 1)]
        for i, p in enumerate(props[:4]):
            ex, ey, side = slots[i]
            r = img_ratio(p)
            sz = sizes[i]
            w, h = (sz, sz / r) if r >= 1 else (sz * r, sz)
            place_image(s, p, ex - w * 0.42 if side < 0 else ex - w * 0.58, ey + (cm(2.6) - h) / 2, w=w,
                        rot=(-8, 6, 8, -6)[i], shadow=True)
    add_source(s, spec.get("source"))


BUILDERS = {"skus": s_skus, "cover": s_cover, "product": s_product, "bullets": s_bullets, "arrows": s_arrows,
            "gallery": s_gallery, "free": s_free, "end": s_end, "stats": s_stats, "compare": s_compare,
            "bars": s_bars, "media": s_media, "article": s_article}


def yt_id(url):
    import re as _re
    m = _re.search(r"(?:v=|youtu\.be/|embed/)([A-Za-z0-9_-]{11})", url or "")
    return m.group(1) if m else None


def yt_thumb(url, cache_dir):
    """Czysta miniatura filmu YouTube (maxresdefault, bez nakładek odtwarzacza) -> plik w cache_dir albo None."""
    vid = yt_id(url)
    if not vid:
        return None
    fn = os.path.join(cache_dir, "yt_%s.jpg" % vid)
    if not os.path.exists(fn):
        import urllib.request
        try:
            os.makedirs(cache_dir, exist_ok=True)
            for q in ("maxresdefault", "hqdefault"):
                try:
                    data = urllib.request.urlopen("https://img.youtube.com/vi/%s/%s.jpg" % (vid, q), timeout=15).read()
                    if len(data) > 5000:
                        open(fn, "wb").write(data)
                        break
                except OSError:
                    continue
        except OSError:
            return None
    return fn if os.path.exists(fn) else None


LINK_HINT = "Kliknij - film otworzy się na YouTube"


def link_overlay(slide, x, y, w, h, url, green="0F763E"):
    """Przycisk play na środku miniatury + czarna pigułka 'Kliknij - film otworzy się na YouTube'; wszystko z tym
    samym hiperłączem (w pokazie 1 klik, w edycji Ctrl+klik). 29.09: YouTube w PowerPoint daje błąd 153, więc
    film otwieramy w przeglądarce - i mówimy o tym wprost na slajdzie."""
    d = min(w, h) * 0.24
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int(x + (w - d) / 2)), Emu(int(y + (h - d) / 2)), Emu(int(d)),
                               Emu(int(d)))
    c.fill.solid(); c.fill.fore_color.rgb = RGBColor.from_string("FFFFFF"); c.line.fill.background()
    c.shadow.inherit = False
    c.name = "Film - przycisk play (link)"
    tw, th = d * 0.36, d * 0.42
    t = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Emu(int(x + w / 2 - tw / 2 + d * 0.05)),
                               Emu(int(y + h / 2 - th / 2)), Emu(int(tw)), Emu(int(th)))
    t.rotation = 90
    t.fill.solid(); t.fill.fore_color.rgb = RGBColor.from_string(green); t.line.fill.background()
    t.shadow.inherit = False
    t.name = "Film - trojkat play (link)"
    ph = cm(0.95)
    pw = min(w - cm(1.0), text_w_emu(LINK_HINT, 13, "Lato Bold") + cm(1.3))
    px, py = x + (w - pw) / 2, y + h - ph - cm(0.45)
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(px)), Emu(int(py)), Emu(int(pw)), Emu(int(ph)))
    pill.adjustments[0] = 0.5
    pill.fill.solid(); pill.fill.fore_color.rgb = RGBColor.from_string("1F1F1F"); pill.line.fill.background()
    pill.shadow.inherit = False
    pill.name = "Film - podpis (link)"
    tf = pill.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = LINK_HINT; r.font.size = Pt(13); r.font.bold = True; r.font.name = "+mn-lt"
    r.font.color.rgb = RGBColor.from_string("FFFFFF")
    for sh in (c, t, pill):
        sh.click_action.hyperlink.address = url
    return pill


def online_videos(out):
    """Znaczniki "ONLINE|adres|promień" -> prawdziwe wideo YouTube w slajdzie (PowerPoint, online_video.ps1).
    Lekcja 29.09.2026: obraz z hiperłączem to NIE film. Zwraca liczbę podmienionych filmów."""
    import subprocess
    from pptx import Presentation as _P
    if not any(sh.name.startswith("ONLINE|") for sl in _P(out).slides for sh in sl.shapes):
        return 0
    ps = os.path.join(os.path.dirname(os.path.abspath(__file__)), "online_video.ps1")
    r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", ps, "-Src", out],
                       capture_output=True, text=True, encoding="cp852", errors="replace")  # konsola PS 5.1 = OEM
    msg = r.stdout.strip() or r.stderr.strip()
    print(msg.encode(sys.stdout.encoding or "utf-8", "replace").decode(sys.stdout.encoding or "utf-8"))
    if r.returncode != 0 or "OK:" not in r.stdout:
        raise SystemExit("BŁĄD: nie wstawiono wideo online - plik ma tylko znacznik zamiast filmu")
    return 1


def main():
    spec_path = sys.argv[1]
    with open(spec_path, encoding="utf-8") as f:
        spec = json.load(f)
    base = os.path.dirname(os.path.abspath(spec_path))

    def resolve(o):  # sciezki obrazow wzgledem pliku spec
        if isinstance(o, dict):
            return {k: (os.path.join(base, v) if (k in ("path", "claim_image", "poster", "video") and isinstance(v, str)
                                                  and not os.path.isabs(v))
                        else [os.path.join(base, q) if isinstance(q, str) and not os.path.isabs(q) else resolve(q) for q in v]
                        if k == "props" and isinstance(v, list) else resolve(v)) for k, v in o.items()}
        if isinstance(o, list):
            return [resolve(v) for v in o]
        return o

    spec = resolve(spec)
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(base, spec.get("output", "prezentacja.pptx"))
    deck = Deck()
    for sl in spec["slides"]:
        n0 = len(deck.prs.slides)
        BUILDERS[sl["type"]](deck, sl)
        if sl.get("morph") is True and len(deck.prs.slides) > n0:
            add_morph(deck.prs.slides[n0])
    out = deck.finish(out)
    online_videos(out)
    print("OK", out, len(spec["slides"]), "slajdow")


if __name__ == "__main__":
    main()
