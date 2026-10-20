# -*- coding: utf-8 -*-
"""Testy usterek ukladu K1-K7 (13.10.2026): urwane haslo, zdublowana liczba-bohater, pusta karta obok pelnej, pusty dol,
sierota / przepelnienie / lamanie hasla z liczba, logo obcej marki na okladce, wykres skumulowany.

    python -m pytest -p no:cacheprovider test_uklad.py -q

Sztuczne slajdy i obrazy (PIL w katalogu tymczasowym); bez plikow usera, bez PowerPointa, bez modeli AI.
"""
import os
import sys

import pytest
from PIL import Image
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_dk as B  # noqa: E402
import szablonuj as S  # noqa: E402


def _png(tmp_path, nazwa, w, h, kolor=(200, 60, 60)):
    p = str(tmp_path / nazwa)
    Image.new("RGB", (w, h), kolor).save(p)
    return p


def _karty(*zdania, **kw):
    """Slajd `uklad` jak z pptx_convert: jeden wiersz kart, w kazdej karcie jeden akapit `p`."""
    row = {"k": "cols", "items": [{"blocks": [{"k": "p", "t": z}]} for z in zdania], "card": True, "kont": True,
           "valign": "t", "pad": 1.1}
    row.update(kw)
    return {"type": "uklad", "title": "Test", "rows": [row]}


def _hasla(sl):
    return [b for r in sl["rows"] if r["k"] == "cols" for it in r["items"] for b in it["blocks"] if b["k"] == "lead"]


# --- K2: liczba-bohater nie powtarza sie w hasle (tylko, gdy po zdjeciu zostaje domknieta fraza) -----------------------
def test_k2_przyimek_przechodzi_do_hero_gdy_fraza_domknieta():
    t = "Wydatki na słone przekąski spadły w tej grupie o 10,1%."
    assert S.bez_liczby((t, None), "10,1%", zdejmij=False) == ("Wydatki na słone przekąski spadły w tej grupie.", None, "o 10,1%")
    assert S.bez_liczby(("GUS prognozuje wzrost do 30%", None), "30%", zdejmij=False) == ("GUS prognozuje wzrost", None, "do 30%")
    assert S.bez_liczby(("Zaburzenia lipidowe dotyczą ponad 50%", None), "50%", zdejmij=False)[2] == "ponad 50%"
    # przecinek po liczbie to zmiana interpunkcji: tylko przy `zdejmij`
    t2 = "Wydatki Polaków na beauty wzrosły w 2025 r. o 7,8%,"
    assert S.bez_liczby((t2, None), "7,8%", zdejmij=False) is None
    assert S.bez_liczby((t2, None), "7,8%", zdejmij=True) == ("Wydatki Polaków na beauty wzrosły w 2025 r.", None, "o 7,8%")


def test_k2_czasownik_zaimek_spojnik_zakres_srodek_zostaja_z_liczba():
    for t, l in [("W 2024 r. osoby 65+ to 7,7 mln,", "7,7 mln"),  # „to” przed liczbą: hasło urwałoby się na „to”
                 ("Skoro numerem 1 jest ciągle energia a grupa stale aktywna sportowo wynosi 23%", "23%"),
                 ("Udział osób 80+ ma wzrosnąć z 4,2% do 11,6%.", "11,6%"),  # „z X do Y”
                 ("GUS prognozuje ok. 30% populacji 65+ w 2060 r.", "30%"),  # w środku zdania
                 ("Cukrzyca dotyczy 21–22 mln osób w Polsce i rośnie", "21–22 mln")]:
        for zd in (False, True):
            assert S.bez_liczby((t, None), l, zdejmij=zd) is None, (t, zd)
    assert S.bez_liczby(("Rynek rośnie szybciej niż i o 12%", None), "12%", zdejmij=True) is None  # ostatnie słowo „i” - spójnik


def test_k2_poczatek_z_myslnikiem_tylko_przy_zdejmowaniu_znakow():
    t = "80% - grupa 1 jak powyżej"
    assert S.bez_liczby((t, None), "80%", zdejmij=False) is None  # zdjęcie myślnika zmieniłoby interpunkcję
    assert S.bez_liczby((t, None), "80%", zdejmij=True) == ("Grupa 1 jak powyżej", None, "80%")
    assert S.bez_liczby(("80% grupa 1 jak powyżej", None), "80%", zdejmij=True) is None  # bez łącznika hasło zostaje z liczbą


def test_k2_wyr_idzie_razem_ze_slowami():
    t, w = "Wydatki Polaków na beauty wzrosły w 2025 r. o 7,8%,", "aa-----------"[:11]
    r = S.bez_liczby((t, w), "7,8%", zdejmij=True)
    assert r[0] == "Wydatki Polaków na beauty wzrosły w 2025 r." and r[1] == "aa---------"[:len(r[0].split())]
    assert len(r[1]) == len(r[0].split())


@pytest.mark.parametrize("znaki", [False, True])
def test_k2_stosuj_slowa_nie_gina(monkeypatch, znaki):
    monkeypatch.setattr(S, "ZDEJMIJ_ZNAKI", znaki)
    sp = _karty("Skoro numerem 1 jest ciągle energia a grupa stale aktywna sportowo wynosi 23%",
                "Wydatki na słone przekąski spadły w tej grupie o 10,1%.",
                "GUS prognozuje ok. 30% populacji 65+ w 2060 r.")
    nowy, dec = S.stosuj(sp, 1)
    assert nowy is not None
    assert S.slowa(sp) == S.slowa(nowy)  # każde słowo (także to w hero z `zast`) jest na slajdzie
    items = nowy["rows"][0]["items"]
    # 1. karta: czasownik przed liczbą - hasło zostaje z liczbą, hero się powtarza
    assert [b["t"] for b in items[0]["blocks"] if b["k"] == "hero"] == ["23%"]
    assert "23%" in [b for b in items[0]["blocks"] if b["k"] == "lead"][0]["t"]
    # 2. karta: „o 10,1%” przechodzi do hero, hasło domknięte kropką
    hero1 = [b for b in items[1]["blocks"] if b["k"] == "hero"][0]
    assert hero1["t"] == "o 10,1%" and hero1.get("zast")
    assert [b for b in items[1]["blocks"] if b["k"] == "lead"][0]["t"].endswith("w tej grupie.")
    # 3. karta: liczba w środku zdania zostaje
    assert not [b for b in items[2]["blocks"] if b["k"] == "hero" and b.get("zast")]


def test_k2_hero_w_liczbach_kontroli_wiernosci(monkeypatch):
    """Liczba z hasła stoi w hero: kontrola liczb (cyfry po wszystkich polach slajdu) widzi ją raz."""
    import re
    monkeypatch.setattr(S, "ZDEJMIJ_ZNAKI", True)
    sp = _karty("Wydatki Polaków na beauty wzrosły w 2025 r. o 7,8%,", "Rynek rośnie szybciej niż rok temu i dalej",
                "Trzeci akapit bez liczby, ale też z hasłem")
    nowy, _ = S.stosuj(sp, 1)
    teksty = [b.get("t", "") for it in nowy["rows"][0]["items"] for b in it["blocks"]]
    assert sorted(re.findall(r"[0-9]+(?:[.,][0-9]+)*", " ".join(teksty))) == ["2025", "7,8"]


def test_k2_dokument_karta_reczna_uzywa_bez_liczby():
    """dokument.py buduje hero sam: liczba znika z hasła tylko, gdy zostaje domknięta fraza (80% - grupa... / o 10,1%)."""
    import dokument as D
    bl = D._karta_reczna(S, {"k": "p", "t": ["80% - grupa 1 jak powyżej"]})["blocks"]
    assert bl[0] == {"k": "hero", "t": "80%", "zast": True} and bl[1]["t"] == "Grupa 1 jak powyżej"
    bl = D._karta_reczna(S, {"k": "p", "t": ["Wydatki na słone przekąski spadły w tej grupie o 10,1%."]})["blocks"]
    assert bl[0]["t"] == "o 10,1%" and bl[0].get("zast") and bl[1]["t"].endswith("w tej grupie.")
    bl = D._karta_reczna(S, {"k": "p", "t": ["Skoro numerem 1 jest ciągle energia a grupa stale aktywna sportowo wynosi 23%"]})["blocks"]
    assert bl[0] == {"k": "hero", "t": "23%"} and bl[1]["t"].endswith("wynosi 23%")  # czasownik przed liczbą: hasło zostaje z liczbą


# --- K1: urwane hasło / pas --------------------------------------------------------------------------------------------
def test_k1_koniec_pasa_i_hasla(monkeypatch):
    rows = [{"k": "band", "t": "Białko –", "fill": "brand"},
            {"k": "cols", "items": [{"blocks": [{"k": "lead", "t": "WHO (2023) zaleca warunkowo,"},
                                                {"k": "lead", "t": "Sarkopenia dotyczy ok. 10% po 80.,"}]}]}]
    monkeypatch.setattr(S, "ZDEJMIJ_ZNAKI", False)
    assert S._k1(rows) == 0 and rows[0]["t"] == "Białko –"
    monkeypatch.setattr(S, "ZDEJMIJ_ZNAKI", True)
    assert S._k1(rows) == 3
    assert rows[0]["t"] == "Białko"
    assert [b["t"] for b in rows[1]["items"][0]["blocks"]] == ["WHO (2023) zaleca warunkowo", "Sarkopenia dotyczy ok. 10% po 80."]


def test_k1_slowa_ignoruja_tylko_brzegowa_interpunkcje():
    a = {"title": "T", "rows": [{"k": "band", "t": "Białko –"}]}
    b = {"title": "T", "rows": [{"k": "band", "t": "Białko"}]}
    c = {"title": "T", "rows": [{"k": "band", "t": "Białko inne"}]}
    assert S.slowa(a) == S.slowa(b) and S.slowa(a) != S.slowa(c)


def test_k12_zdjete_znaki_zapisane_w_slajdzie(monkeypatch):
    """Zdjęty znak interpunkcji jest deklarowany w slajdzie (`zdjete_znaki`) - kontrola wierności może go odjąć od „znaki”."""
    sp = _karty("Wydatki Polaków na beauty wzrosły w 2025 r. o 7,8%, a najszybciej rośnie grupa", "Rynek rośnie szybciej niż rok temu i dalej",
                "Trzeci akapit bez liczby, ale też z hasłem")
    monkeypatch.setattr(S, "ZDEJMIJ_ZNAKI", True)
    assert S.stosuj(sp, 1)[0]["zdjete_znaki"] == ","
    monkeypatch.setattr(S, "ZDEJMIJ_ZNAKI", False)
    assert "zdjete_znaki" not in S.stosuj(sp, 1)[0]


# --- K3: krótka karta w rzędzie kart ---------------------------------------------------------------------------------
def _rzad_z_krotka_karta():
    dluga = "Żelki funkcjonalne to najszybciej rosnący format, szczególnie wśród Gen Z i millenialsów, a liderem sprzedaży są napoje."
    row = {"k": "cols", "card": True, "kont": True, "valign": "t", "pad": 1.1, "rwane_ok": True, "items": [
        {"blocks": [{"k": "h", "t": "Wzrost napędza kolagen."}, {"k": "p", "t": dluga, "max_pt": 18}]},
        {"blocks": [{"k": "lead", "t": "Komunikacja od anti-agingu do prewencji", "color": "ink"}]},
        {"blocks": [{"k": "h", "t": "Wzrost napędza kolagen."}, {"k": "p", "t": dluga, "max_pt": 18}]}]}
    return {"type": "uklad", "title": "Dane wejściowe", "rows": [row], "valign": "t"}


def test_k3_krotka_karta_ma_tresc_wysrodkowana():
    sp = _rzad_z_krotka_karta()
    ms, _f = B.uklad_layout(sp)
    assert ms[0]["vc"] == [False, True, False]
    # jawny valign karty wygrywa
    sp["rows"][0]["items"][1]["valign"] = "m"
    assert B.uklad_layout(sp)[0][0]["vc"][1] is False  # `valign: m` rysuje się tak samo, bez znacznika vc


def test_k3_ocena_nie_karze_wysrodkowanej_karty():
    sp = _rzad_z_krotka_karta()
    kody = [u["kod"] for u in S.ocena(sp, B)["usterki"]]
    assert "KARTA_PUSTA" not in kody


def test_k3_rysowanie_nie_wywala_i_tekst_krotkiej_karty_jest_w_srodku(tmp_path):
    sp = _rzad_z_krotka_karta()
    B.build({"theme": "shop", "slides": [sp]}, str(tmp_path / "k3.pptx"))
    s = Presentation(B.build.last_out).slides[0]
    karty = sorted((sh for sh in s.shapes if sh.shape_type == 1 and sh.height > B.cm(5)), key=lambda x: x.left)
    assert len(karty) == 3
    sr = karty[1]
    teksty = [sh for sh in s.shapes if sh.has_text_frame and "Komunikacja" in sh.text_frame.text]
    assert teksty
    t = teksty[0]
    gora, dol = t.top - sr.top, (sr.top + sr.height) - (t.top + t.height)
    assert abs(gora - dol) < 0.25 * sr.height and gora > 0.15 * sr.height  # blok mniej więcej na środku, nie przy górze


# --- K4: pusty dół ------------------------------------------------------------------------------------------------------
def test_k4_obrazy_ulozone_tak_by_byly_wieksze(tmp_path):
    items = [{"image": _png(tmp_path, "a.png", 1200, 300), "caption": ""}, {"image": _png(tmp_path, "b.png", 1200, 300), "caption": ""}]
    width, h = B.W - 2 * B.MX, B.cm(11)
    rzad = B._pics_rzad(items, width, h, 0)
    najl = B._pics_najlepszy(items, width, h, 0)
    pole = lambda u: sum(w * hh for _i, _x, w, hh, _y in u)
    assert pole(najl) >= 1.25 * pole(rzad)  # dwa szerokie obrazy: jeden pod drugim jest większy niż obok siebie
    assert len({round(y) for _i, _x, _w, _h, y in najl}) == 2
    assert all(hh <= h + 1 for _i, _x, _w, hh, _y in najl)


def test_k4_obrazy_wysokie_zostaja_w_jednym_rzedzie(tmp_path):
    items = [{"image": _png(tmp_path, "c.png", 300, 600), "caption": ""}, {"image": _png(tmp_path, "d.png", 300, 600), "caption": ""}]
    width, h = B.W - 2 * B.MX, B.cm(11)
    assert B._pics_najlepszy(items, width, h, 0) == B._pics_rzad(items, width, h, 0)


def test_k4_obrazy_w_pliku_nie_sa_rozciagane(tmp_path):
    items = [{"image": _png(tmp_path, "e.png", 1200, 300), "caption": ""}, {"image": _png(tmp_path, "f.png", 1200, 300), "caption": ""}]
    sp = {"type": "uklad", "title": "Wykresy", "rows": [{"k": "pics", "items": items, "min_h": 4.0, "h": 8.5}], "valign": "t"}
    B.build({"theme": "shop", "slides": [sp]}, str(tmp_path / "k4.pptx"))
    s = Presentation(B.build.last_out).slides[0]
    pics = [sh for sh in s.shapes if sh.shape_type == 13 and sh.width > B.cm(5) and sh.height > B.cm(2)]
    assert len(pics) == 2
    for p in pics:
        assert abs(p.width / p.height - 4.0) < 0.02  # proporcje 4:1 jak plik, bez rozciągania i przycinania


def _slajd_z_karta(min_h):
    row = {"k": "cols", "items": [{"blocks": [{"k": "lead", "t": "Krótka teza slajdu", "color": "ink"}]}], "card": True,
           "kont": True, "valign": "t", "pad": 1.1, "min_h": min_h}
    return {"type": "uklad", "title": "Test", "rows": [row], "valign": "t"}


def test_k4_wysrodkuj_blok_gdy_dol_pusty_ponad_35_procent():
    assert S._wysrodkuj(_slajd_z_karta(6.5), B)["valign"] == "m"  # ~46% pola pustego, pas nad i pod < 18% wysokości slajdu
    assert S._wysrodkuj(_slajd_z_karta(10.5), B).get("valign") == "t"  # pole prawie pełne
    assert S._wysrodkuj(_slajd_z_karta(3.0), B).get("valign") == "t"  # treść pływałaby (PLYWAJACA_TRESC): zostaje u góry


def test_k4_krotka_tabela_wyzsze_wiersze_i_srodek_pola(tmp_path):
    def tabela(ile):
        return {"type": "table", "title": "Cennik", "head": ["Produkt", "Cena"],
                "rows": [["Pozycja %d" % i, "%d zł" % (10 + i)] for i in range(ile)]}
    B.build({"theme": "shop", "slides": [tabela(3)]}, str(tmp_path / "t3.pptx"))
    kr = Presentation(B.build.last_out).slides[0]
    nag = [sh for sh in kr.shapes if sh.has_text_frame and sh.text_frame.text.strip() == "Produkt"][0]
    assert nag.top > B.CT + B.cm(1.0) and nag.text_frame.paragraphs[0].runs[0].font.size.pt >= 15
    B.build({"theme": "shop", "slides": [tabela(12)]}, str(tmp_path / "t12.pptx"))
    dl = Presentation(B.build.last_out).slides[0]
    nag = [sh for sh in dl.shapes if sh.has_text_frame and sh.text_frame.text.strip() == "Produkt"][0]
    assert nag.top == B.CT and nag.text_frame.paragraphs[0].runs[0].font.size.pt == 13  # długa tabela jak dotąd


# --- K5: sierota, przepełnienie, łamanie hasła z liczbą -----------------------------------------------------------------
def test_k5_rowne_lamanie_usuwa_sierote():
    t = "Zadbać o to, aby komunikacja była satysfakcjonującą."
    ls = B.display_lines(t, 26, B.cm(9))
    assert not B._orphan(ls) and " ".join(ls) == t
    r = B._lamanie_rowne("Marka zdrowa i smaczna dla każdej osoby", 24, B.cm(8.5), 3)
    assert r is None or (len(r) == 3 and len(r[-1].split()) >= 2)


def test_k5_haslo_z_liczba_i_jednostka_nie_lamie_sie(tmp_path):
    row = {"k": "cols", "card": True, "kont": True, "valign": "t", "pad": 1.1, "rwane_ok": True, "items": [
        {"blocks": [{"k": "lead", "t": "300 g wody", "color": "ink"}]},
        {"blocks": [{"k": "lead", "t": "Zdrowie jelit zaczyna się od mikrobioty i błonnika", "color": "ink"}]},
        {"blocks": [{"k": "lead", "t": "Błonnik pomaga trawieniu każdego dnia", "color": "ink"}]},
        {"blocks": [{"k": "lead", "t": "Cztery karty obok siebie w wąskich kolumnach", "color": "ink"}]}]}
    sp = {"type": "uklad", "title": "Test", "rows": [row], "valign": "t"}
    ms, _f = B.uklad_layout(sp)
    lead = [d for kind, _y, _h, d in ms[0]["cols"][0][0] if kind == "txt"][0]
    assert lead["lines"] == ["300 g wody"] and lead["pt"] >= B.HASLO_MIN_PT
    assert B._krotkie_z_liczba("300 g wody") and not B._krotkie_z_liczba("Zdrowie jelit zaczyna się od mikrobioty")


def test_k5_pole_tekstu_ma_zapas_na_zaokraglony_skok_linii(tmp_path):
    """PowerPoint zaokrągla skok linii do pełnego punktu (27,54 -> 28): 5 linii x 28 = 140 pt, nie 137,7 (verify: PRZEPELNIENIE +2pt)."""
    sp = {"type": "uklad", "title": "T", "rows": [{"k": "lead", "t": "Projekt –\nRedesign znaku\nmarki i visuali\nposzczególnych\nopakowań",
                                                   "pt": 25.5}], "valign": "t"}
    B.build({"theme": "shop", "slides": [sp]}, str(tmp_path / "k5.pptx"))
    s = Presentation(B.build.last_out).slides[0]
    tb = [sh for sh in s.shapes if sh.has_text_frame and "Redesign" in sh.text_frame.text][0]
    pt = tb.text_frame.paragraphs[0].line_spacing.pt
    assert tb.height >= (len(tb.text_frame.paragraphs) * round(pt) + 1) * 12700 - 1


def test_k5_ochrona_przed_sierota_nie_robi_linii_szerszej_od_pola():
    """Słowo schodzi do ostatniej linii tylko, gdy ta się zmieści („też zapotrzebowanie” 122,9 pt w polu 120,5 pt łamało się w PowerPoincie)."""
    w = 120.5 * B.bd.EMU_PT if hasattr(B.bd, "EMU_PT") else 120.5 * 12700
    meas = B.FONT_REF["body"][1]
    for l in B.lines_for("a rośnie też zapotrzebowanie", 13.5, w, "body"):
        assert B.bd.text_w_emu(l.replace(" ", " "), 13.5, meas) <= w, l


def test_k5_sklejony_token_szerszy_od_pola_zmniejsza_pismo_bloku():
    """„i mikroskładniki.” to jeden niełamliwy token (zawieszka): pismo bloku maleje, aż się zmieści - nie linia poza polem."""
    w = 141.5 * 12700
    t = "Smak i mikroskładniki."
    assert B._uk_wrap("display", t, 20.5, w) is None  # przy 20,5 pt token (145 pt) jest szerszy od pola
    pt = B._slowo_pt("display", t, 20.5, w, 20.5)
    assert pt < 20.5 and B.bd.text_w_emu("i mikroskładniki.", pt, "Mindset") <= w * 0.96
    assert B._uk_wrap("display", t, pt, w) is not None


# --- K6: logo obcej marki na okładce ------------------------------------------------------------------------------------
def _okladka(tmp_path, v, logo=True):
    sp = {"type": "cover_text", "variant": v, "kicker": "Raport", "title": "Strategia produktowa na rok 2027", "subtitle": "Podtytuł",
          "meta": "Autor, 13.10.2026"}
    if logo:
        sp["logo"] = _png(tmp_path, "logo.png", 800, 300, (30, 90, 200))
    B.build({"theme": "shop", "slides": [sp]}, str(tmp_path / ("c%d%s.pptx" % (v, "l" if logo else "n"))))
    return Presentation(B.build.last_out).slides[0]


@pytest.mark.parametrize("v", [1, 2, 3])
def test_k6_logo_obce_proporcje_wysokosc_i_brak_kolizji(tmp_path, v):
    s = _okladka(tmp_path, v)
    lg = [sh for sh in s.shapes if sh.name == "!!logo-obce"]
    assert len(lg) == 1
    lg = lg[0]
    assert lg.height <= B.cm(3.5) + 1 and lg.rotation == 0
    assert abs(lg.width / lg.height - 800 / 300) < 0.02
    assert lg.left > B.W / 2 - 1  # prawa część okładki
    for sh in s.shapes:
        if sh is lg or not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        x = min(lg.left + lg.width, sh.left + sh.width) - max(lg.left, sh.left)
        y = min(lg.top + lg.height, sh.top + sh.height) - max(lg.top, sh.top)
        assert not (x > 0 and y > 0), "logo zachodzi na %r" % sh.text_frame.text[:30]


def test_k6_bez_klucza_logo_okladka_jak_dotad(tmp_path):
    for v in (1, 2, 3):
        s = _okladka(tmp_path, v, logo=False)
        assert not [sh for sh in s.shapes if sh.name == "!!logo-obce"]
    sp = {"type": "cover_text", "variant": 1, "title": "T", "logo": str(tmp_path / "brak.png")}
    B.build({"theme": "shop", "slides": [sp]}, str(tmp_path / "brak.pptx"))  # plik logo nie istnieje: okładka bez logo, bez wyjątku


# --- K7: wykres skumulowany -------------------------------------------------------------------------------------------
@pytest.mark.parametrize("klucze,grupowanie,nakladka", [({}, "clustered", "-10"), ({"stacked": True}, "stacked", "100"),
                                                         ({"percent": True}, "percentStacked", "100")])
def test_k7_typ_wykresu_w_xml(tmp_path, klucze, grupowanie, nakladka):
    sp = {"type": "columns", "title": "Udział", "categories": ["2023", "2024"],
          "series": [{"name": "A", "values": [60, 40]}, {"name": "B", "values": [40, 60]}]}
    sp.update(klucze)
    B.build({"theme": "shop", "slides": [sp]}, str(tmp_path / "k7.pptx"))
    s = Presentation(B.build.last_out).slides[0]
    ch = [sh for sh in s.shapes if getattr(sh, "has_chart", False) and sh.has_chart][0].chart
    xml = ch._chartSpace.xml
    assert 'c:grouping val="%s"' % grupowanie in xml
    assert 'c:overlap val="%s"' % nakladka in xml
    if klucze:
        assert "outEnd" not in xml  # OUTSIDE_END w kolumnach skumulowanych to błąd PowerPointa
