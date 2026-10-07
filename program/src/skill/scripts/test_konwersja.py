# -*- coding: utf-8 -*-
"""Testy konwersji 1:1 (pptx_convert): pojemniki z szablonu dobierane do ksztaltu tresci (08.10.2026).

    python -m pytest test_konwersja.py -q

Czesc 1 - jednostkowe, na sztucznych wierszach (bez plikow). Czesc 2 - integracyjne na prawdziwych prezentacjach
(pomijane, gdy pliku nie ma na tym komputerze). Zadne z nich nie uruchamia PowerPointa ani modeli AI.
"""
import collections
import os
import sys

import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_dk  # noqa: E402
import pptx_convert as pc  # noqa: E402

PREZ = r"D:\Marketing\- POLSKA\02 - FIRMOWE MATERIAŁY\PREZENTACJE"
PLIKI = {
    "co_dalej": os.path.join(PREZ, "06.10.2026 - strategia", "DK_co_dalej_update.pptx"),
    "datesy": os.path.join(PREZ, "DATESY", "DOBRA KALORIA - DATESY.pptx"),
    "kulki": os.path.join(PREZ, "24.09.2026 - KULKI z kreatyną", "DK_KULKI z kreatyną.pptx"),
    "onlien": os.path.join(PREZ, "onlien.pptx"),
}
# Znane usterki wiernosci PRZED zmiana (07.10): konwersja nie moze ich pomnozyc
USTERKI_PRZED = {"co_dalej": 0, "datesy": 0, "kulki": 1, "onlien": 152}


# --- 1. jednostkowe ---------------------------------------------------------------------------------------------
def _w(*lines):
    return [{"t": t, "pt": 36, "li": False, "pusty": False} for t in lines]


def test_zdanie_progi_slow():
    n12 = " ".join(["slowo"] * 12)
    n13 = " ".join(["slowo"] * 13)
    assert not pc._zdanie_grupa(_w(n12))
    assert pc._zdanie_grupa(_w(n13))
    assert not pc._zdanie_grupa(_w(" ".join(["slowo"] * 7) + "."))
    assert pc._zdanie_grupa(_w(" ".join(["slowo"] * 8) + "."))
    assert not pc._zdanie_grupa(_w(" ".join(["slowo"] * 8) + "?"))  # pytanie bywa haslem


def _row(k, t, pole, pt=36.0, duze=True, **kw):
    r = {"k": k, "t": t, "_pole": pole, "_pt": pt, "_duze": duze}
    r.update(kw)
    return r


def test_dobra_zla_wiadomosc_dwie_karty():
    out = pc._pojemniki([
        _row("h", "Dobra wiadomość – rynek zdrowych przekąsek rośnie", 1, _fl="iiiiiii"),
        _row("lead", "Zła wiadomość – rośniemy wolniej\nniż sam rynek", 1, _fl="aaaaa----"[:8]),
    ])
    assert len(out) == 1 and out[0]["k"] == "cols" and out[0]["kont"] and len(out[0]["items"]) == 2
    h = [it["blocks"][0] for it in out[0]["items"]]
    assert [b["k"] for b in h] == ["h", "h"]
    assert h[0]["t"] == "Dobra wiadomość –" and h[1]["t"] == "Zła wiadomość –"  # myślnik z oryginału zostaje (09.10)


def test_cytat_i_karta_hasla():
    out = pc._pojemniki([
        _row("lead", "Brands grow through constant exposure\nto new buyers (Byron Sharp, Ehrenberg-Bass).", 1,
             _fl="--aa-aa--------"[:11]),
        _row("lead", "Produkt: jasno zdefiniowany", 2, pt=36),
        _row("lead", "Potrzeba, grupa docelowa, okazja", 2, pt=47),
        _row("lead", "= Rozwiązanie produktowe", 3, pt=36),
    ])
    assert [r["k"] for r in out] == ["quote", "cols"]
    assert out[0]["author"] == "(Byron Sharp, Ehrenberg-Bass)." and out[0]["t"].endswith("new buyers")  # nawiasy i kropka zostają
    assert len(out[1]["items"]) == 1 and len(out[1]["items"][0]["blocks"]) == 3  # R6 odrzucone (pt 47/36, '=')


def test_trzy_fakty_i_wniosek_goly():
    pole_a, pole_b = 1, 2
    out = pc._pojemniki([
        _row("p", " ".join(["fakt"] * 18) + ".", pole_a, pt=32), _row("p", " ".join(["dane"] * 9) + ".", pole_a, pt=32),
        _row("p", " ".join(["liczba"] * 9) + ".", pole_a, pt=32),
        _row("lead", "Ważna kwestia – to nie jest jednolity segment", pole_b, pt=30),
    ])
    assert [r["k"] for r in out] == ["cols", "lead"]
    assert len(out[0]["items"]) == 3 and out[1]["t"].startswith("Ważna")


def test_etykieta_ct_ta_zostaje_goly_h():
    out = pc._pojemniki([_row("lead", "CT: młodzi z segmentu 20-35 lat\nTA: mężczyźni 35-55 lat", 1)])
    assert out[0]["k"] == "h" and out[0]["align"] == "l" and not out[0].get("kont")


def test_wstep_przed_kafelkami_i_akapit_w_karcie():
    out = pc._pojemniki([
        _row("h", "Kierunek zmian:", 1),
        {"k": "chips", "items": ["A", "B", "C"], "per_row": 3},
        _row("p", " ".join(["slowo"] * 25) + ".", 2, pt=18, duze=False),
    ])
    assert [r["k"] for r in out] == ["h", "chips", "cols"]
    assert out[2]["items"][0]["blocks"][0]["k"] == "p"


def test_lista_pozycji_laczy_oraz_i_nie_skleja_sierot():
    assert pc._pozycje_listy(["Ewaluacja portfolio", "Kierunki", "Kreacja nowych segmentów", "oraz kategorii"], 6) == \
        ["Ewaluacja portfolio", "Kierunki", "Kreacja nowych segmentów oraz kategorii"]
    assert pc._pozycje_listy(["Polifenole i flawonoidy", "Beta-karoten", "Cynk", "Magnez"], 4) is not None
    assert pc._pozycje_listy(["Jedno zdanie.", "Drugie zdanie.", "Trzecie zdanie."], 6) is None


# --- 2. integracyjne --------------------------------------------------------------------------------------------
_CACHE = {}


def _zbuduj(nazwa, tmp_path_factory):
    if nazwa in _CACHE:
        return _CACHE[nazwa]
    plik = PLIKI[nazwa]
    if not os.path.isfile(plik):
        pytest.skip("brak pliku %s" % plik)
    rob = str(tmp_path_factory.mktemp("rob_" + nazwa))
    model = pc.extract(plik, rob)
    spec = pc.compose(model, {"wizualizacje": False, "wyjscie": rob})
    out = os.path.join(rob, "wynik.pptx")
    build_dk.build(spec, out)
    pptx = build_dk.build.last_out
    cov = pc.coverage(plik, pptx, spec["mapa"], rob=rob)
    wier = pc.wiernosc(model, pptx, spec["mapa"])
    _CACHE[nazwa] = (model, spec, cov, wier)
    return _CACHE[nazwa]


@pytest.mark.parametrize("nazwa", sorted(PLIKI))
def test_tresc_i_wiernosc(nazwa, tmp_path_factory):
    model, spec, cov, wier = _zbuduj(nazwa, tmp_path_factory)
    assert cov["brakuje"] == []
    assert len(wier["usterki"]) <= USTERKI_PRZED[nazwa]
    if nazwa != "onlien":  # onlien (instrukcja szablonu) dzieli dlugie slajdy juz od 06.10
        assert len(spec["slides"]) == len(model["slajdy"])


@pytest.mark.parametrize("nazwa", sorted(PLIKI))
def test_tlo_papierowe_i_pismo_miesci_sie(nazwa, tmp_path_factory):
    _model, spec, _cov, _wier = _zbuduj(nazwa, tmp_path_factory)
    for sp in spec["slides"]:
        assert sp.get("bg") != "card"  # krem tylko jako karta, nigdy tlo calego slajdu
        if sp["type"] == "uklad" and sp.get("rows"):
            assert pc._miesci(sp)


def test_co_dalej_haslo_to_nie_zdanie_i_kazdy_slajd_ma_pojemnik(tmp_path_factory):
    model, spec, _cov, _wier = _zbuduj("co_dalej", tmp_path_factory)
    for i, sp in enumerate(spec["slides"], 1):
        if sp["type"] != "uklad":
            continue
        for r in sp["rows"]:
            blocks = [b for it in r.get("items", []) for b in it.get("blocks", [])] if r["k"] == "cols" else [r]
            for b in blocks:
                if b["k"] in ("lead", "h", "sub", "big") and b.get("t"):
                    assert len(b["t"].split()) < 13, "slajd %d: Mindset z %d slowami" % (i, len(b["t"].split()))
    for src in range(6, 43):
        idx = spec["mapa"][str(src)]
        sps = [spec["slides"][k] for k in idx]
        if all(sp["type"] != "uklad" for sp in sps):
            continue  # przerywnik
        assert any(r["k"] in ("cols", "quote", "chips", "pics") or r.get("li") for sp in sps for r in sp["rows"]), \
            "slajd %d bez pojemnika" % src


# --- 3. runda 2 (08.10): hiperlacza zamiast 'Link:', wielkosc liter, zdanie z kropka to nie haslo -------------------
def test_zdaniowo_mala_litera_w_srodku_zdania():
    t = "Matcha to nie tylko smak ale pewien Manifest – młodzi, nie chcą przepłacać Za dobra luksusowe ale Chcą się"
    out = pc._zdaniowo([t], set(), {"za", "chcą", "smak"})[0]
    assert " za dobra " in out and " chcą się" in out and "Manifest" in out  # 'Manifest' nie ma malej wersji w pliku
    assert pc._zdaniowo(["To jest. Za chwile"], set(), {"za"})[0].startswith("To jest. Za")  # po kropce zostaje


def test_hiperlacza_bez_napisu_link(tmp_path_factory):
    model, spec, _cov, _wier = _zbuduj("co_dalej", tmp_path_factory)
    assert "Link:" not in str(spec["slides"])
    lk = {sp["zrodlo"]: sp["links"] for sp in spec["slides"] if sp.get("links")}
    assert lk[7][0]["t"] == "Byron Sharp" and lk[41][0]["t"].startswith("Bezkaloryczny")


def test_hiperlacze_w_pliku_na_tych_samych_slowach(tmp_path_factory):
    from pptx import Presentation
    _model, spec, _cov, _wier = _zbuduj("co_dalej", tmp_path_factory)
    rob = str(tmp_path_factory.mktemp("lk"))
    build_dk.build(spec, os.path.join(rob, "w.pptx"))
    prs = Presentation(build_dk.build.last_out)
    idx = spec["mapa"]["7"][0]
    runs = [r for sh in prs.slides[idx].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs
            if r.hyperlink and r.hyperlink.address]
    assert [r.text for r in runs] == ["Byron Sharp"]


def test_zdanie_z_kropka_nie_jest_mindsetem(tmp_path_factory):
    _model, spec, _cov, _wier = _zbuduj("co_dalej", tmp_path_factory)
    for sp in spec["slides"]:
        for r in sp.get("rows", []):
            blocks = [b for it in r.get("items", []) for b in it.get("blocks", [])] if r["k"] == "cols" else [r]
            for b in blocks:
                if b.get("k") in ("lead", "h", "big") and b.get("t"):  # 'sub' = wniosek pod obrazami (R10a, K10)
                    assert not (len(b["t"].split()) >= 8 and b["t"].rstrip().endswith((".", "!"))), b["t"]


def test_komentarz_po_kartach_nie_wiekszy_niz_karty(tmp_path_factory):
    _model, spec, _cov, _wier = _zbuduj("co_dalej", tmp_path_factory)
    sp = next(s for s in spec["slides"] if s["zrodlo"] == 37)
    ms, _f = build_dk.uklad_layout(sp)
    karta = max(d["pt"] for parts, _w in ms[0]["cols"] for (kind, _y, _h, d) in parts if kind == "txt")
    kom = [m["pt"] for m, r in zip(ms, sp["rows"]) if r["k"] == "p"]
    assert kom and max(kom) <= karta


# --- 4. runda 3 (09.10): znak w znak, karty wg tresci, obrazy do brzegow, zakonczenie na srodku -----------------------
def test_interpunkcja_bez_roznic_co_dalej(tmp_path_factory):
    _model, _spec, _cov, wier = _zbuduj("co_dalej", tmp_path_factory)
    assert wier["znaki"] == [], wier["znaki"]  # zadnego zdjetego nawiasu, myslnika ani kropki (s06, s07, s09)


def test_s09_kafle_zielone_bez_wiodacego_myslnika(tmp_path_factory):
    _model, spec, _cov, wier = _zbuduj("co_dalej", tmp_path_factory)
    ch = [r for k in spec["mapa"]["9"] for r in spec["slides"][k]["rows"] if r["k"] == "chips"][0]
    # 10.10 (runda 4): myslnik listy to marker (jak punktor): kafle go nie maja, zawsze zielen marki (nie akcent)
    assert len(ch["items"]) == 4 and not any(t.startswith(("–", "-")) for t in ch["items"]) and ch.get("fill", "brand") == "brand"
    assert wier["znaki"] == [] and any(m["slajd"] == 9 for m in wier["markery"])  # zdjety marker wpisany w pokrycie


def test_wiernosc_zlapie_zdjety_nawias():
    assert pc._interp("(Byron Sharp, Ehrenberg-Bass).") - pc._interp("Byron Sharp, Ehrenberg-Bass") ==         pc.collections.Counter({"(": 1, ")": 1, ".": 1})


def test_obrazy_w_calosci_proporcje_bez_kadrowania(tmp_path_factory):
    import io
    from PIL import Image
    from pptx import Presentation
    _model, spec, _cov, _wier = _zbuduj("co_dalej", tmp_path_factory)
    rob = str(tmp_path_factory.mktemp("img"))
    build_dk.build(spec, os.path.join(rob, "w.pptx"))
    prs = Presentation(build_dk.build.last_out)
    for src in (15, 25, 34, 35, 36, 38, 42):
        for k in spec["mapa"][str(src)]:
            for sh in prs.slides[k].shapes:
                if sh.shape_type == 13 and not sh.name.startswith("!!"):
                    w, h = Image.open(io.BytesIO(sh.image.blob)).size
                    assert abs(sh.width / sh.height - w / h) / (w / h) < 0.02, "slajd %d: obraz rozciagniety" % src
                    assert not (sh.crop_left or sh.crop_right or sh.crop_top or sh.crop_bottom), "slajd %d: kadrowanie" % src


def test_s34_rzad_obrazow_od_brzegu_do_brzegu(tmp_path_factory):
    from pptx import Presentation
    _model, spec, _cov, _wier = _zbuduj("co_dalej", tmp_path_factory)
    rob = str(tmp_path_factory.mktemp("rz"))
    build_dk.build(spec, os.path.join(rob, "w.pptx"))
    prs = Presentation(build_dk.build.last_out)
    pics = [sh for sh in prs.slides[spec["mapa"]["34"][0]].shapes if sh.shape_type == 13 and not sh.name.startswith("!!")]
    assert abs(min(p.left for p in pics) - build_dk.MX) < 20000  # < 0,6 mm
    assert abs(max(p.left + p.width for p in pics) - (build_dk.W - build_dk.MX)) < 20000


def test_s22_dwie_karty_rownej_wysokosci_tresc_od_gory(tmp_path_factory):
    _model, spec, _cov, _wier = _zbuduj("co_dalej", tmp_path_factory)
    sp = spec["slides"][spec["mapa"]["22"][0]]
    ms, _f = build_dk.uklad_layout(sp)
    m = next(x for x in ms if x["kind"] == "cols")
    assert len({round(x) for x in m["hs"]}) == 1 and m["valign"] == "t"  # 10.10: równa wysokość kart, krótsza treść od góry
    assert min(m["chs"]) < 0.85 * max(m["chs"])  # treść faktycznie nierówna
    assert not any("łatwy" in ln and "do przegryzienia" not in ln and ln.endswith("łatwy")
                   for parts, _w in m["cols"] for (kind, _y, _h, d) in parts if kind == "txt" for ln in d["lines"])
    for n in (18, 24, 37, 40):  # trzy karty / podobna tresc: wszystkie karty rowne
        ms, _f = build_dk.uklad_layout(spec["slides"][spec["mapa"][str(n)][0]])
        assert all(len({round(x) for x in m["hs"]}) == 1 for m in ms if m["kind"] == "cols")


def test_zakonczenie_na_srodku_pionowo(tmp_path_factory):
    from pptx import Presentation
    _model, spec, _cov, _wier = _zbuduj("co_dalej", tmp_path_factory)
    sp = spec["slides"][spec["mapa"]["43"][0]]
    assert sp["type"] == "end" and sp["valign"] == "m"
    rob = str(tmp_path_factory.mktemp("end"))
    build_dk.build(dict(spec, slides=[sp]), os.path.join(rob, "e.pptx"))
    sl = Presentation(build_dk.build.last_out).slides[0]
    logo = next(sh for sh in sl.shapes if sh.name == "!!logo")
    tag = next(sh for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.startswith("#"))
    srodek = (logo.top + tag.top + tag.height - build_dk.cm(0.2)) / 2  # tekst hasztagu zajmuje ok. 1,4 z 1,8 cm
    assert abs(srodek - build_dk.H / 2) < build_dk.cm(0.8)


# --- 4a. runda 4 (10.10): karty na pelna szerokosc, wielkosc liter po lamaniu, cytat jak s20, kafle, podpis pod rzedem ----
def _linie_kart(spec, n):
    """[(liczba kart w wierszu, [wiersze tekstu Lato])] dla slajdu zrodla n (z miar ukladu)."""
    out = []
    for k in spec["mapa"][str(n)]:
        sp = spec["slides"][k]
        ms, _f = build_dk.uklad_layout(sp)
        for m in ms:
            if m["kind"] == "cols":
                out.append((len(m["cols"]), m["pad"], [d for parts, _w in m["cols"] for (kind, _y, _h, d) in parts
                                                      if kind == "txt" and d["font"] == "body"]))
    return out


def test_s18_s21_s40_linie_w_kartach_nie_sa_rwane(tmp_path_factory):
    _model, spec, _cov, _wier = _zbuduj("co_dalej", tmp_path_factory)
    for n in (18, 21, 40):
        for ncols, pad, bloki in _linie_kart(spec, n):
            assert pad <= build_dk.cm(0.61)  # pole tekstu = szerokosc karty minus 2 x ~0,6 cm
            for d in bloki:
                slow = sum(len(x.split()) for x in d["lines"])
                if slow >= 8:  # 3 wąskie karty obok siebie: >= 3,5 słowa w wierszu; karty jedna pod drugą: pełna szerokość
                    assert slow / len(d["lines"]) >= (3.5 if ncols >= 3 else 4), (n, d["lines"])


def test_kontynuacja_mala_litera_i_nowy_akapit_wielka(tmp_path_factory):
    _model, spec, _cov, _wier = _zbuduj("co_dalej", tmp_path_factory)
    t26 = " ".join(bl["t"] for r in spec["slides"][spec["mapa"]["26"][0]]["rows"] for it in r.get("items", [])
                   for bl in it["blocks"])
    assert "ważny ale równie" in t26 and "Ale równie" not in t26  # wiersz-kontynuacja po wierszu bez kropki: mała litera
    bl41 = [bl["t"] for r in spec["slides"][spec["mapa"]["41"][0]]["rows"] for it in r.get("items", []) for bl in it["blocks"]]
    assert any(t.startswith("Dla segmentu szkolnego") for t in bl41)  # po łączu kończącym zdanie: osobny akapit, wielka litera
    assert not any("udaru dla segmentu" in t for t in bl41)


def test_split_vt_lacze_konczy_zdanie():
    t = "Erytrytol bez kropki na koncu\vDla segmentu szkolnego"
    assert [x for x, _h, _e in pc._split_vt(t, True)] == ["Erytrytol bez kropki na koncu Dla segmentu szkolnego"]
    assert [x for x, _h, _e in pc._split_vt(t, True, [True, False])] == ["Erytrytol bez kropki na koncu", "Dla segmentu szkolnego"]
    assert [x for x, _h, _e in pc._split_vt("Ala ma\vkota", True, [True, False])] == ["Ala ma kota"]  # po łączu, ale mała litera


def test_chronione_slowo_po_kropce_to_poczatek_zdania(tmp_path_factory):
    model, _spec, _cov, _wier = _zbuduj("co_dalej", tmp_path_factory)
    assert "Ale" not in pc._chronione(model)  # „…fizycznej. Ale samo białko” - początek zdania, nie nazwa własna


def test_s41_jeden_rozmiar_pisma_min_14(tmp_path_factory):
    _model, spec, _cov, _wier = _zbuduj("co_dalej", tmp_path_factory)
    bloki = _linie_kart(spec, 41)[0][2]
    assert len({d["pt"] for d in bloki}) == 1 and bloki[0]["pt"] >= 14 and len(bloki) == 3  # WHO + Erytrytol + Dla segmentu


def test_cytat_jak_s20_i_autor_pogrubiony(tmp_path_factory):
    from pptx import Presentation
    _model, spec, _cov, _wier = _zbuduj("co_dalej", tmp_path_factory)
    rob = str(tmp_path_factory.mktemp("cyt"))
    build_dk.build(spec, os.path.join(rob, "w.pptx"))
    sl = Presentation(build_dk.build.last_out).slides[spec["mapa"]["7"][0]]
    runy = [(r, p.text) for sh in sl.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs]
    q = next(r for r, t in runy if t.startswith("Brands grow") or t.startswith("new buyers") or "constant" in t)
    assert q.font.size.pt >= 28 and q.font.italic  # s20 szablonu: 30 pt kursywa
    aut = [r for r, t in runy if t in ("(", "Byron Sharp", ", Ehrenberg-Bass).") or t.endswith("Ehrenberg-Bass).")]
    assert aut and all(r.font.bold and r.font.size.pt >= 14 for r in aut)  # autor pogrubiony Lato >= 14 pt (hiperłącze zostaje)


def test_kafle_rowna_siatka():
    m = build_dk._uk_chips({"k": "chips", "items": ["– A", "Dluzsze haslo wiersza", "B", "Haslo srednie"], "per_row": 2},
                           1.0, build_dk.W - 2 * build_dk.MX)
    assert m["ws"][0] == m["ws"][2] and m["ws"][1] == m["ws"][3]  # kafle jednej kolumny tej samej szerokosci
    assert m["fill"] == "brand"


def test_s38_podpis_pod_srodkiem_rzedu_obrazow(tmp_path_factory):
    from pptx import Presentation
    _model, spec, _cov, _wier = _zbuduj("co_dalej", tmp_path_factory)
    rob = str(tmp_path_factory.mktemp("sub"))
    build_dk.build(spec, os.path.join(rob, "w.pptx"))
    sl = Presentation(build_dk.build.last_out).slides[spec["mapa"]["38"][0]]
    pics = [sh for sh in sl.shapes if sh.shape_type == 13 and not sh.name.startswith("!!")]
    sub = next(sh for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.startswith("A konkurencja"))
    cx = (min(p.left for p in pics) + max(p.left + p.width for p in pics)) / 2
    assert abs((sub.left + sub.width / 2) - cx) < build_dk.cm(0.1)


# --- 5. tabela regresji (4 prezentacje z zadania; bez PowerPointa i modeli AI): pytest test_konwersja.py -k tabela -s -
REG = os.path.join(r"C:\Users\KRZYSZ~1.WIE\AppData\Local\Temp\claude\d--Marketing---POLSKA-02---FIRMOWE-MATERIA-Y-PREZENTACJE"
                   r"\724cc04a-b3dc-4f72-b73d-790b5bc9f790\scratchpad\trojki\v5\r4\regresja")
REG_PLIKI = {
    "co_dalej": PLIKI["co_dalej"],
    "onlien": PLIKI["onlien"],
    "DATESY": PLIKI["datesy"],
    "KULKI_stary": os.path.join(PREZ, "24.09.2026 - KULKI z kreatyną", "KULKI z kreatyną - stary styl.pptx"),
}
_GOLE = ("lead", "h", "p")


def _gola(sp):
    return sp["type"] == "uklad" and bool(sp.get("rows")) and all(r["k"] in _GOLE and not r.get("li") for r in sp["rows"])


@pytest.mark.parametrize("nazwa", sorted(REG_PLIKI))
def test_tabela_regresji(nazwa):
    plik = REG_PLIKI[nazwa]
    if not os.path.isfile(plik):
        pytest.skip("brak pliku %s" % plik)
    rob = os.path.join(REG, nazwa)
    os.makedirs(rob, exist_ok=True)
    model = pc.extract(plik, rob)
    spec = pc.compose(model, {"wizualizacje": False, "wyjscie": rob})
    build_dk.build(spec, os.path.join(rob, "wynik.pptx"))
    pptx = build_dk.build.last_out
    cov = pc.coverage(plik, pptx, spec["mapa"], rob=rob)
    wier = pc.wiernosc(model, pptx, spec["mapa"])
    sl = spec["slides"]
    uk = [s for s in sl if s["type"] == "uklad"]
    typy = collections.Counter(s["type"] for s in sl)
    gole = sum(_gola(s) for s in uk)
    wiersze = collections.Counter(r["k"] for s in uk for r in s.get("rows", []))
    wiersz = "%s | slajdy %d->%d | tresc %s | wiernosc usterek %d, znaki %d | uklad %d, szablon %d, gole %d (%.0f%% ukladow) | %s" % (
        nazwa, len(model["slajdy"]), len(sl), "brakuje %d" % len(cov["brakuje"]), len(wier["usterki"]),
        len(wier.get("znaki", [])), len(uk), len(sl) - len(uk), gole, 100.0 * gole / max(1, len(uk)),
        dict(typy))
    with open(os.path.join(REG, "tabela.txt"), "a", encoding="utf-8") as f:
        f.write(wiersz + "\n  wiersze: %s\n  usterki: %s\n" % (dict(wiersze), wier["usterki"][:10]))
    print(wiersz)
    assert cov["brakuje"] == [], cov["brakuje"][:5]
    assert wier["usterki"] == [] or len(wier["usterki"]) <= USTERKI_PRZED.get(nazwa, 0), wier["usterki"][:5]
    assert wier.get("znaki", []) == [], wier["znaki"][:5]
    if nazwa != "onlien":  # onlien (instrukcja szablonu) dzieli dlugie slajdy od 06.10
        assert len(sl) == len(model["slajdy"])
