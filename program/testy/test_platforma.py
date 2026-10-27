# -*- coding: utf-8 -*-
"""Testy 1.3 (macOS): logika wersji na Maca bez prawdziwego Maca. Podmieniamy platforma.mac() / platforma.windows() i atrapy
subprocess (pbcopy, pbpaste, open, osascript). Budowa prezentacji idzie na prawdziwym silniku, ale z zakazem uruchamiania
PowerShella i skryptów .ps1 (na Macu ich nie ma).
    python -m pytest -q -p no:cacheprovider WORK\\logs\\test_platforma.py

Katalog roboczy testów: %TEMP%\\claude-tmp\\2026-10-27-prezentacje-mac\\testy (testy niczego nie kasują - sprząta właściciel sesji).
Testy Windowsa (prawdziwy schowek, rejestr) są pomijane poza Windows."""
import os
import shlex
import shutil
import subprocess
import sys
import tempfile
import types

import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
WORK = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(WORK, "src", "app"))
import paths  # noqa: E402

_KAND = [os.path.join(WORK, "src", "skill"), os.path.join(os.path.expanduser("~"), ".claude", "skills", "prezentacje")]
for _k in [k for k in _KAND if os.path.exists(os.path.join(k, "scripts", "scenariusz.py")) and os.path.exists(os.path.join(k, "scripts", "dokument.py"))]:
    paths.SKILL_DIR, paths.SKILL_SCRIPTS = _k, os.path.join(_k, "scripts")
    paths.VERIFY_PS = os.path.join(_k, "scripts", "verify.ps1")
    break
import engine  # noqa: E402
import folder_prez  # noqa: E402
import gui  # noqa: E402
import platforma  # noqa: E402
import update  # noqa: E402
import szybka  # noqa: E402

ROOT = os.path.join(tempfile.gettempdir(), "claude-tmp", "2026-10-27-prezentacje-mac", "testy")
WZOR = os.path.join(paths.SKILL_DIR, "assets", "scenariusz-wzor.md")
POLSKIE = "Zażółć gęślą jaźń - Prezentacje Dobra Kaloria"
tylko_windows = pytest.mark.skipif(sys.platform != "win32", reason="prawdziwe API Windows")


@pytest.fixture
def tmp():
    os.makedirs(ROOT, exist_ok=True)
    return tempfile.mkdtemp(prefix="t-", dir=ROOT)


@pytest.fixture
def mac(monkeypatch):
    """Program myśli, że działa na macOS."""
    monkeypatch.setattr(platforma, "mac", lambda: True)
    monkeypatch.setattr(platforma, "windows", lambda: False)


@pytest.fixture
def dom(tmp, monkeypatch):
    """Podmieniony katalog domowy (HOME i USERPROFILE), żeby ~/Applications, ~/.local/bin i ~/Documents były w katalogu testowym."""
    d = os.path.join(tmp, "home")
    os.makedirs(d)
    monkeypatch.setenv("HOME", d)
    monkeypatch.setenv("USERPROFILE", d)
    monkeypatch.setattr(os.path, "expanduser", lambda p: d + p[1:] if p.startswith("~") else p)
    return d


class Run:
    """Atrapa subprocess.run: zapisuje wywołania, zwraca zadany kod i wyjście."""
    def __init__(self, code=0, stdout=b"", wyjatek=None):
        self.code, self.stdout, self.wyjatek, self.calls = code, stdout, wyjatek, []

    def __call__(self, cmd, *a, **kw):
        self.calls.append((list(cmd), kw))
        if self.wyjatek:
            raise self.wyjatek
        return types.SimpleNamespace(returncode=self.code, stdout=self.stdout, stderr=b"")


# --- schowek (pbcopy / pbpaste) ------------------------------------------------------------------------------------------------
def test_schowek_ustaw_mac_pbcopy_utf8(mac, monkeypatch):
    run = Run()
    monkeypatch.setattr(platforma.subprocess, "run", run)
    tekst = "# SCENARIUSZ DK\n" + POLSKIE
    assert platforma.schowek_ustaw(tekst) is True
    cmd, kw = run.calls[0]
    assert cmd == ["pbcopy"] and kw["input"] == tekst.encode("utf-8")
    assert kw["env"]["LANG"] == "en_US.UTF-8" and kw["env"]["LC_ALL"] == "en_US.UTF-8"     # bez LANG pbcopy zapisałby MacRoman
    assert "shell" not in kw
    assert gui.set_clipboard("x") is True                                                    # okno woła tę samą warstwę


def test_schowek_ustaw_mac_bledy(mac, monkeypatch):
    monkeypatch.setattr(platforma.subprocess, "run", Run(code=1))
    assert platforma.schowek_ustaw("x") is False
    monkeypatch.setattr(platforma.subprocess, "run", Run(wyjatek=FileNotFoundError("pbcopy")))
    assert platforma.schowek_ustaw("x") is False
    monkeypatch.setattr(platforma.subprocess, "run", Run(wyjatek=subprocess.TimeoutExpired("pbcopy", 10)))
    assert platforma.schowek_ustaw("x") is False


def test_schowek_pobierz_mac_pbpaste(mac, monkeypatch):
    run = Run(stdout=("# SCENARIUSZ DK\n" + POLSKIE).encode("utf-8"))
    monkeypatch.setattr(platforma.subprocess, "run", run)
    assert platforma.schowek_pobierz() == "# SCENARIUSZ DK\n" + POLSKIE
    cmd, kw = run.calls[0]
    assert cmd == ["pbpaste"] and kw["env"]["LANG"] == "en_US.UTF-8"
    assert gui.get_clipboard() == "# SCENARIUSZ DK\n" + POLSKIE
    assert folder_prez.wyglada_na_scenariusz(platforma.schowek_pobierz())
    for zly in (Run(stdout=b""), Run(code=1, stdout=b"x"), Run(stdout=b"x" * (platforma.LIMIT_SCHOWKA + 1)),
                Run(wyjatek=FileNotFoundError("pbpaste")), Run(wyjatek=subprocess.TimeoutExpired("pbpaste", 3))):
        monkeypatch.setattr(platforma.subprocess, "run", zly)
        assert platforma.schowek_pobierz() is None


@tylko_windows
def test_schowek_windows_bez_zmian():
    stary = platforma.schowek_pobierz()
    try:
        assert platforma.schowek_ustaw("# SCENARIUSZ DK\n" + POLSKIE)
        assert platforma.schowek_pobierz() == "# SCENARIUSZ DK\n" + POLSKIE
    finally:
        platforma.schowek_ustaw(stary or "")


# --- otwieranie plików, folderów, adresów, aplikacji -------------------------------------------------------------------------------
def test_otworz_mac_to_open(mac, monkeypatch):
    run = Run()
    monkeypatch.setattr(platforma.subprocess, "run", run)
    platforma.otworz("/Users/x/Documents/Prezentacje Dobra Kaloria/Źródła")
    platforma.otworz("https://claude.ai/new")
    platforma.otworz("claude://cowork/new?q=a%20b&folder=%2FUsers%2Fx")
    assert [c[0] for c in run.calls] == [["open", "/Users/x/Documents/Prezentacje Dobra Kaloria/Źródła"], ["open", "https://claude.ai/new"],
                                         ["open", "claude://cowork/new?q=a%20b&folder=%2FUsers%2Fx"]]
    monkeypatch.setattr(platforma.subprocess, "run", Run(code=1))
    with pytest.raises(OSError):      # jak os.startfile: wywołujący łapie OSError
        platforma.otworz("claude://nic")


def test_pokaz_w_folderze_mac_open_R(mac, monkeypatch):
    zapis = []
    monkeypatch.setattr(platforma.subprocess, "Popen", lambda cmd, *a, **k: zapis.append(cmd))
    platforma.pokaz_w_folderze("/Users/x/Prezentacje/Oferta Żółć.pptx")
    platforma.uruchom_aplikacje("/Applications/Claude.app")
    assert zapis == [["open", "-R", "/Users/x/Prezentacje/Oferta Żółć.pptx"], ["open", "/Applications/Claude.app"]]


def test_okno_mac_otwiera_pliki_i_foldery(mac, tmp, monkeypatch):
    f = os.path.join(tmp, "Oferta.pptx")
    open(f, "wb").write(b"x")
    run, popen = Run(), []
    monkeypatch.setattr(platforma.subprocess, "run", run)
    monkeypatch.setattr(platforma.subprocess, "Popen", lambda cmd, *a, **k: popen.append(cmd))
    api = gui.Api()
    assert api.open_path(f) == {"ok": True, "komunikat": ""} and run.calls[-1][0] == ["open", f]
    assert api.open_folder(f)["ok"] and popen[-1] == ["open", "-R", f]
    assert api.open_folder(tmp)["ok"] and run.calls[-1][0] == ["open", tmp]
    monkeypatch.setattr(platforma.subprocess, "run", Run(code=1))
    r = api.open_path(f)
    assert r["ok"] is False and "Nie mogę otworzyć" in r["komunikat"]
    r = api.open_folder(tmp)
    assert r["ok"] is False and "Finderze" in r["komunikat"] and "Eksploratorze" not in r["komunikat"]


def test_znajdz_aplikacje_mac(mac, dom):
    assert platforma.znajdz_aplikacje("Claude") is None and platforma.schemat_claude() is False
    os.makedirs(os.path.join(dom, "Applications", "Claude.app"))
    assert platforma.znajdz_aplikacje("Claude") == os.path.join(dom, "Applications", "Claude.app")
    assert platforma.schemat_claude() is True and engine.claude_protokol() is True and engine.find_app("Claude").endswith("Claude.app")
    for zla in ("", None, "Cla;ude", "../Claude", "Claude'; rm"):
        assert platforma.znajdz_aplikacje(zla) is None


def test_czat_w_aplikacji_i_w_przegladarce_na_macu(mac, dom, tmp, monkeypatch):
    run, popen = Run(), []
    monkeypatch.setattr(platforma.subprocess, "run", run)
    monkeypatch.setattr(platforma.subprocess, "Popen", lambda cmd, *a, **k: popen.append(cmd))
    monkeypatch.setattr(gui, "set_clipboard", lambda t: True)
    api = gui.Api()
    r = api._otworz_czat("claude", "polecenie")                     # brak aplikacji: przeglądarka
    assert r["ok"] and r["gdzie"] == "przegladarka" and run.calls[-1][0] == ["open", "https://claude.ai/new"]
    os.makedirs(os.path.join(dom, "Applications", "Claude.app"))
    r = api._otworz_czat("claude", "polecenie")                     # jest Claude.app
    assert r["ok"] and r["gdzie"] == "aplikacja" and popen[-1] == ["open", os.path.join(dom, "Applications", "Claude.app")]
    monkeypatch.setattr(platforma.subprocess, "Popen", lambda *a, **k: (_ for _ in ()).throw(OSError("brak")))
    r = api._otworz_czat("claude", "polecenie")
    assert r["ok"] is False and "Nie udało się otworzyć czatu Claude" in r["blad"]


# --- Claude Code w Terminalu (osascript) ---------------------------------------------------------------------------------------------
def test_znajdz_claude_na_macu_szuka_poza_path(mac, dom, monkeypatch):
    monkeypatch.setattr(platforma.shutil, "which", lambda n: None)
    assert platforma.znajdz_claude() is None
    os.makedirs(os.path.join(dom, ".local", "bin"))
    k = os.path.join(dom, ".local", "bin", "claude")
    open(k, "w").write("#!/bin/sh\n")
    assert platforma.znajdz_claude() == k          # aplikacja z Findera nie ma PATH z terminala
    monkeypatch.setattr(platforma.shutil, "which", lambda n: "/opt/homebrew/bin/claude")
    assert platforma.znajdz_claude() == "/opt/homebrew/bin/claude"


@pytest.mark.parametrize("nazwa", ["Prezentacje Dobra Kaloria", POLSKIE, "O'Neil & $HOME `x` (kopia)", "bez-spacji"])
def test_claude_code_osascript_cytowanie(mac, tmp, monkeypatch, nazwa):
    folder = os.path.join(tmp, nazwa)
    os.makedirs(folder)
    claude = "/Users/Zażółć/.local/bin/claude"
    monkeypatch.setattr(platforma.shutil, "which", lambda n: claude)
    run = Run()
    monkeypatch.setattr(platforma.subprocess, "run", run)
    assert engine.uruchom_claude_code(folder) is True                 # engine.uruchom_claude_code = warstwa platformy
    cmd, kw = run.calls[0]
    assert isinstance(cmd, list) and cmd[0] == "osascript" and "shell" not in kw
    linia = cmd[-1]
    # jedna warstwa cytowania: powłoka w Terminalu dostaje dokładnie: cd <folder> && <claude> <polecenie>
    assert shlex.split(linia) == ["cd", folder, "&&", claude, engine.POLECENIE_CLAUDE_CODE]
    skrypt = [cmd[i + 1] for i, x in enumerate(cmd[:-1]) if x == "-e"]
    assert cmd[1:-1] == [x for e in skrypt for x in ("-e", e)] and 'tell application "Terminal"' in skrypt
    assert "do script (item 1 of argv)" in skrypt and "activate" in skrypt
    assert not any(folder in e or engine.POLECENIE_CLAUDE_CODE in e for e in skrypt)   # ścieżka nie trafia do tekstu AppleScript


def test_claude_code_osascript_porazka(mac, tmp, monkeypatch):
    monkeypatch.setattr(platforma.shutil, "which", lambda n: "/usr/local/bin/claude")
    monkeypatch.setattr(platforma.subprocess, "run", Run(code=1))          # np. brak zgody na sterowanie Terminalem
    assert platforma.uruchom_claude_code(tmp, "x") is False
    monkeypatch.setattr(platforma.subprocess, "run", Run(wyjatek=subprocess.TimeoutExpired("osascript", 30)))
    assert platforma.uruchom_claude_code(tmp, "x") is False
    monkeypatch.setattr(platforma.subprocess, "run", Run(wyjatek=FileNotFoundError("osascript")))
    assert platforma.uruchom_claude_code(tmp, "x") is False
    run = Run()
    monkeypatch.setattr(platforma.subprocess, "run", run)
    assert platforma.uruchom_claude_code(os.path.join(tmp, "nie ma"), "x") is False and not run.calls   # brak folderu
    monkeypatch.setattr(platforma.shutil, "which", lambda n: None)
    monkeypatch.setattr(platforma, "znajdz_claude", lambda: None)
    assert platforma.uruchom_claude_code(tmp, "x") is False and not run.calls                         # brak `claude`


def test_ai_cele_na_macu(mac, dom, tmp, monkeypatch):
    monkeypatch.setattr(platforma.shutil, "which", lambda n: None)
    monkeypatch.setattr(folder_prez, "ustawienia", lambda: {})
    assert [c["id"] for c in engine.ai_cele()] == ["claude", "chatgpt", "gemini"]
    os.makedirs(os.path.join(dom, "Applications", "Claude.app"))
    assert [c["id"] for c in engine.ai_cele()] == ["claude", "claude_cowork", "chatgpt", "gemini", "claude_code"]


# --- ścieżki i baza folderu prezentacji --------------------------------------------------------------------------------------------
def test_katalog_uzytkownika_mac_i_windows(mac, dom, monkeypatch):
    assert platforma.katalog_uzytkownika() == os.path.join(dom, "Library", "Application Support", "Dobra Kaloria", "Stworz prezentacje")
    monkeypatch.setattr(platforma, "mac", lambda: False)
    monkeypatch.setattr(platforma, "windows", lambda: True)
    monkeypatch.setenv("LOCALAPPDATA", os.path.join(dom, "AppData"))
    assert platforma.katalog_uzytkownika() == os.path.join(dom, "AppData", "Dobra Kaloria", "Stworz prezentacje")


def test_baza_prezentacji_na_macu_to_dokumenty(mac, dom, monkeypatch):
    monkeypatch.delenv("DK_PREZ_BAZA", raising=False)
    monkeypatch.setattr(folder_prez, "ustawienia", lambda: {})
    assert folder_prez.baza_m() == ""                                            # dysku M: na Macu nie ma
    baza, m_ok = folder_prez.baza_domyslna()
    assert m_ok is False and baza == os.path.join(dom, "Documents", "Prezentacje Dobra Kaloria") == folder_prez.dokumenty_baza()
    monkeypatch.setenv("DK_PREZ_BAZA", dom)                                      # podmiana (testy, inny dysk) działa jak na Windows
    assert folder_prez.baza_m() == dom and folder_prez.baza_domyslna() == (dom, True)


def test_analiza_na_macu_nie_ostrzega_o_dysku_m(mac, dom, tmp, monkeypatch):
    monkeypatch.delenv("DK_PREZ_BAZA", raising=False)
    monkeypatch.setattr(folder_prez, "ustawienia", lambda: {})
    monkeypatch.setattr(paths, "USER_DIR", os.path.join(tmp, "user"))
    engine._dcache.clear()
    p = os.path.join(tmp, "pobrane", "opis.txt")
    os.makedirs(os.path.dirname(p))
    open(p, "w", encoding="utf-8").write("# Tytuł\n\n" + "To jest zdanie, które ma kilka słów i kończy się kropką. " * 30)
    r = engine.analyze_many([p])
    assert r["ok"] and r["tryb"] == "pakiet" and r["m_dostepny"] is False
    assert r["baza"] == folder_prez.dokumenty_baza() and not any("Dysk M:" in u for u in r["uwagi"])


# --- silnik bez PowerPointa i bez .ps1 -------------------------------------------------------------------------------------------------
def test_powerpoint_na_macu_niedostepny(mac):
    assert platforma.powerpoint_dostepny() is False and engine.powerpoint_available() is False
    assert gui.Api().get_info()["powerpoint"] is False and gui.Api().get_info()["system"] == "mac"


@tylko_windows
def test_system_windows_w_get_info():
    assert gui.Api().get_info()["system"] == "windows"


@pytest.fixture
def bez_windowsa(mac, tmp, monkeypatch):
    """Silnik jak na Macu: brak PowerPointa, PowerShell i skrypty .ps1 zakazane, katalog użytkownika w testowym."""
    wywolania = []
    run0, popen0 = subprocess.run, subprocess.Popen

    def sprawdz(cmd):
        t = " ".join(str(x) for x in (cmd if isinstance(cmd, (list, tuple)) else [cmd])).lower()
        if "powershell" in t or ".ps1" in t or "taskkill" in t or "explorer" in t:
            wywolania.append(t)
            raise AssertionError("na Macu nie wolno uruchamiać: " + t)
    monkeypatch.setattr(subprocess, "run", lambda cmd, *a, **k: sprawdz(cmd) or run0(cmd, *a, **k))
    monkeypatch.setattr(subprocess, "Popen", lambda cmd, *a, **k: sprawdz(cmd) or popen0(cmd, *a, **k))
    monkeypatch.setattr(engine, "_ps", lambda *a, **k: sprawdz(["_ps"] + list(a)))
    monkeypatch.setattr(szybka, "OCR_PS", "")                       # na Macu szybka.ocr_texts i tak nic nie robi (os.name != "nt")
    user = os.path.join(tmp, "user")
    os.makedirs(user)
    monkeypatch.setattr(paths, "USER_DIR", user)
    monkeypatch.setenv("DK_PREZ_BAZA", os.path.join(tmp, "PREZENTACJE"))
    engine._dcache.clear()
    engine._pcache.clear()
    yield tmp
    assert wywolania == []


NIE_SPRAWDZONO = "Nie sprawdzono w PowerPoint: na Macu program nie robi podglądu slajdów."


def _opis(res):
    assert res["ok"], res.get("blad") or res.get("traceback")
    assert os.path.isfile(res["pptx"])
    uwagi = res.get("uwagi_pelne") or res.get("uwagi")        # folder produktu podaje "uwagi", reszta także "uwagi_pelne"
    assert any(u.startswith(NIE_SPRAWDZONO) for u in uwagi), uwagi
    assert res["qa"]["wykonano"] is False and res["qa"]["png"] is None


def test_budowa_dokumentu_na_macu(bez_windowsa):
    p = os.path.join(bez_windowsa, "pobrane", "opis.txt")
    os.makedirs(os.path.dirname(p))
    open(p, "w", encoding="utf-8").write("# Dobra przekąska\n\n" + "\n\n".join(
        "Akapit numer %d mówi o czymś ważnym dla klienta: kulki są pyszne, lekkie i zrobione z prawdziwych owoców." % i for i in range(9)))
    _opis(engine.safe_build(p, {}))             # bez "bez_qa": idzie ścieżką, która na Windows woła PowerPoint


def test_budowa_ze_scenariusza_na_macu(bez_windowsa):
    baza = os.path.join(bez_windowsa, "PREZENTACJE")
    zr = os.path.join(bez_windowsa, "pobrane", "opis.txt")
    os.makedirs(os.path.dirname(zr))
    open(zr, "w", encoding="utf-8").write("Opis produktu. " * 20)
    from PIL import Image
    jpg = os.path.join(bez_windowsa, "pobrane", "baton-przyklad.jpg")      # obraz, do którego odwołuje się wzór scenariusza
    Image.new("RGB", (900, 600), (40, 140, 70)).save(jpg)
    folder, _ = folder_prez.utworz_folder(baza, "2026-10-27 - Test na Macu", [zr, jpg])
    folder_prez.nowa_wersja_scenariusza(folder, open(WZOR, encoding="utf-8-sig").read())
    res = engine.safe_build(folder, {})
    _opis(res)
    assert res["tryb"] == "scenariusz" and os.path.dirname(res["pptx"]) == folder


def test_konwersja_pptx_na_macu(bez_windowsa):
    from pptx import Presentation
    from pptx.util import Inches
    p = os.path.join(bez_windowsa, "obca", "Oferta obca.pptx")
    os.makedirs(os.path.dirname(p))
    prs = Presentation()
    for t in ("Oferta dla handlu", "Zalety produktu: smak, skład, cena", "Dziękujemy za uwagę"):
        sl = prs.slide_layouts[6]
        s = prs.slides.add_slide(sl)
        s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = t
    prs.save(p)
    res = engine.safe_build(p, {"cel": "wiernie"})
    _opis(res)


PRODUKT = r"D:\Marketing\- POLSKA\02 - FIRMOWE MATERIAŁY\PREZENTACJE\24.09.2026 - KULKI z kreatyną"


@pytest.mark.skipif(not os.path.isdir(PRODUKT), reason="brak wzorcowego folderu produktu na tym komputerze")
def test_budowa_folderu_produktu_na_macu(bez_windowsa):
    prod = os.path.join(bez_windowsa, "KULKI z kreatyną TEST")
    os.makedirs(prod)
    for f in os.listdir(PRODUKT):
        if f.startswith("Karta wprowadzenia") or f.lower().endswith(".docx"):
            shutil.copy2(os.path.join(PRODUKT, f), prod)
    _opis(engine.safe_build(prod, {}))


def test_dokument_doc_na_macu_to_czytelny_komunikat(mac, tmp, monkeypatch):
    import dokument
    monkeypatch.setattr(dokument.sys, "platform", "darwin")
    with pytest.raises(ValueError) as e:
        dokument._przez_word(os.path.join(tmp, "stary.doc"), tmp)
    assert ".doc" in str(e.value) and "Macu" in str(e.value) and "Traceback" not in str(e.value)


def test_wideo_online_i_qa_dokumentu_na_macu_nie_wola_powershella(mac, tmp, monkeypatch):
    import build_deck
    import dokument
    from pptx import Presentation
    monkeypatch.setattr(build_deck.sys, "platform", "darwin")
    monkeypatch.setattr(subprocess, "run", lambda *a, **k: pytest.fail("PowerShell na Macu"))
    p = os.path.join(tmp, "film.pptx")
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    from pptx.util import Inches
    s.shapes.add_picture(os.path.join(paths.SKILL_DIR, "assets", "logo_green_box.png"), Inches(1), Inches(1)).name = "ONLINE|https://youtu.be/x|0"
    prs.save(p)
    assert build_deck.online_videos(p) == 0          # zostaje obraz z linkiem, bez wyjątku
    assert dokument._qa({}, p, tmp, "") is None      # „Na Macu bez podglądu slajdów: otwórz plik i sprawdź.”


# --- aktualizacja ---------------------------------------------------------------------------------------------------------------------
def test_aktualizacja_na_macu_tylko_komunikat_ze_strona(mac, monkeypatch):
    ok, powod = update.can_self_update()
    assert ok is False and update.STRONA_MAC in powod and "Macu" in powod
    with pytest.raises(RuntimeError):
        update.apply("/tmp/x")                      # nic nie podmienia
    rel = {"tag_name": "v9.9.9", "assets": [{"name": "Stworz-prezentacje-9.9.9-Windows.zip", "browser_download_url": "https://github.com/x.zip",
                                                 "size": 1}], "html_url": "https://github.com/inyfinn/Prezentacje-Dobra-Kaloria/releases/tag/v9.9.9"}
    monkeypatch.setattr(update, "_get", lambda url, timeout=8: _Odp(rel))
    r = update.check()
    assert r["jest"] is True and r["mac"] is True and r["sam"] is False and r["strona"] == update.STRONA_MAC
    assert "url" not in r                           # nic do pobrania: Mac dostaje tylko link
    api = gui.Api()
    out = api.check_update()
    assert out["mac"] is True and out["strona"] == update.STRONA_MAC and "url" not in out
    rel["tag_name"] = "v0.0.1"
    assert update.check()["jest"] is False


class _Odp:
    def __init__(self, rel):
        import json
        self.b = json.dumps(rel).encode("utf-8")

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def read(self):
        return self.b


@tylko_windows
def test_aktualizacja_na_windows_bez_zmian(monkeypatch):
    monkeypatch.setattr(paths, "FROZEN", False)
    assert update.can_self_update() == (False, "wersja deweloperska")
    monkeypatch.setattr(paths, "FROZEN", True)
    ok, powod = update.can_self_update()           # zależnie od miejsca programu: dysk wspólny, brak zapisu albo można
    assert "Macu" not in powod and (ok or powod in ("program stoi na dysku wspólnym", "brak prawa zapisu w folderze programu"))
    assert platforma.dysk_lokalny(tempfile.gettempdir()) in (True, False)


# --- okno: filtr plików i start -------------------------------------------------------------------------------------------------------
def test_filtry_okien_wyboru_plikow_przechodza_walidacje_pywebview(monkeypatch):
    """Okno wyboru pliku w pywebview (macOS i Windows) odrzuca opis filtra z przecinkiem (ValueError z parse_file_type)."""
    from webview.util import parse_file_type
    zlapane = []

    class Okno:
        def create_file_dialog(self, tryb, directory="", allow_multiple=False, file_types=()):
            zlapane.extend(file_types)
            return None
    api = gui.Api()
    api._window = Okno()
    api.pick_file()
    api.pick_files()
    assert len(zlapane) == 4
    for f in zlapane:
        parse_file_type(f)


def test_start_okna_mac_bez_wymuszania_gui(mac, monkeypatch):
    zapis = {}
    monkeypatch.setattr(gui.webview, "create_window", lambda *a, **k: types.SimpleNamespace(events=None))
    monkeypatch.setattr(gui.webview, "start", lambda *a, **k: zapis.update(k))
    monkeypatch.setattr(gui.os, "_exit", lambda c: zapis.update(wyjscie=c))
    gui.run()
    assert zapis["gui"] is None and zapis["wyjscie"] == 0
    assert gui._window_geometry() == (None, None, gui.WIN_W, gui.WIN_H, gui.MIN_W, gui.MIN_H)    # bez ctypes.windll


def test_cli_na_macu_nie_podlacza_konsoli_windows(mac, monkeypatch):
    import cli
    monkeypatch.setattr(cli.sys, "platform", "darwin")
    monkeypatch.setattr(cli.sys, "stdout", None)
    cli._attach_console()                           # nie dotyka ctypes.windll
    assert cli.sys.stdout is None


def test_sciezka_cli_w_polecenie_agenta(mac, monkeypatch):
    import scenariusz
    folder = os.path.join(ROOT, "agent")
    monkeypatch.setattr(scenariusz.sys, "platform", "win32")
    assert 'stworz-cli.exe "' in scenariusz.polecenie_agenta(folder)
    monkeypatch.setattr(scenariusz.sys, "platform", "darwin")
    monkeypatch.setattr(scenariusz.sys, "frozen", True, raising=False)
    monkeypatch.setattr(scenariusz.sys, "executable", "/Applications/Stwórz prezentację.app/Contents/MacOS/stworz-prezentacje")
    t = scenariusz.polecenie_agenta(folder)
    assert '"/Applications/Stwórz prezentację.app/Contents/MacOS/stworz-prezentacje" "%s" --scenariusz' % folder in t
    assert "stworz-cli.exe" not in t
